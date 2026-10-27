# SPDX-FileCopyrightText: Copyright (c) 2025-2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""The ground-truth review package: Pending stays Pending, decisions carry forward, P1 pronoun flags, no overwrite."""

from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import gt_candidate_build  # noqa: E402
import gt_review_package as pkg  # noqa: E402
import review_backup  # noqa: E402
from test_ground_truth_correction import FRAMES, _changes, _first, _proposal, setup  # noqa: E402, F401
from test_ground_truth_review import _decide  # noqa: E402

REWRITE = "A person hits another person and pushes him."


@pytest.fixture
def candidate(setup, tmp_path):  # noqa: F811
    source, manifest, prior = setup
    _, _, changes = _changes(source, manifest)
    review = [c for c in changes if c.change_type in ("needs_video_review", "rule_based_proposal")]
    proposals = []
    for change in review:
        if change.rule == "R25" and not any(p["decision"] == "propose_change" for p in proposals):
            proposals.append(_proposal(change, proposed_value=REWRITE))
        elif change.rule == "R22":
            proposals.append(_proposal(change, decision="uncertain", proposed_value=None, evidence=[],
                                       confidence="low"))
        else:
            proposals.append(_proposal(change, decision="keep_original", proposed_value=None))
    out = tmp_path / "review"
    path = out / "ai-proposals-v0.2.jsonl"
    path.write_text("".join(json.dumps(p) + "\n" for p in proposals))
    frames = tmp_path / "frames"
    for incident in {c.incident_id for c in changes}:
        (frames / incident).mkdir(parents=True, exist_ok=True)
        (frames / incident / "frames.json").write_text(json.dumps(FRAMES))
    gt_candidate_build.build(source, manifest, prior, path, out, "v0.2", frames_dir=frames)
    r25 = next(p["change_id"] for p in proposals if p["decision"] == "propose_change")
    (out / "review-overrides-v0.2.json").write_text(json.dumps(
        {"pronoun_fixes": {r25: [["pushes him", "pushes them"]]}, "recommendations": {}}))
    return source, manifest, out / "Ground Truth Candidate v0.2.xlsx", frames, len(review), r25


def _build(candidate, tmp_path, version="v0.2.1", **kwargs):
    source, _, path, frames, _, _ = candidate
    return pkg.build(path, version, source=source, out_dir=path.parent, html_dir=tmp_path / "html", frames_dir=frames,
                     **kwargs)


def _log(path: Path) -> dict:
    ws = openpyxl.load_workbook(path)["Change Log"]
    header = [c.value for c in ws[1]]
    return {row[0]: dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True)}


def test_every_pending_entry_stays_pending_and_auto_verified_rows_are_kept(candidate, tmp_path):
    source, _, path, _, pending, _ = candidate
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, path)}
    result = _build(candidate, tmp_path)
    assert {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, path)} == before
    total = result["checks"]["statuses"]["Pending"]  # the review entries plus any informational ones
    assert total >= pending and result["rows"] == total and result["statuses"] == {"Pending": total}
    log = _log(Path(result["output"]))
    assert sum(r["Review status"] == "Auto-verified" for r in log.values()) == result["checks"]["statuses"][
        "Auto-verified"]
    assert all(r["Reviewer"] is None and r["Accepted value"] is None for r in log.values())
    checks = result["checks"]
    assert checks["pending_applied"] == 0 and checks["pending_cells_changed"] == 0
    assert checks["auto_verified_not_traced"] == 0 and checks["auto_verified_applied_and_passed"]
    assert review_backup.review_state(Path(result["output"]))["decisions"] == 0


def test_p1_flags_violations_with_a_fix_and_leaves_decisions_to_the_reviewer(candidate, tmp_path):
    result = _build(candidate, tmp_path)
    row = _log(Path(result["output"]))[candidate[5]]
    assert "P1 violation" in row["Flags"] and row["Suggested accepted value"].endswith("pushes them.")
    assert row["Review status"] == "Pending" and row["Accepted value"] is None
    assert result["flags"]["UNCERTAIN"] == 1


def test_a_rewrite_with_a_permitted_pronoun_is_flagged_for_confirmation():
    flag = pkg.pronoun_flag("CL-1", "propose_change", "Figure X, which appears male; he waves.", {})
    assert flag.startswith("PRONOUN: permitted under P1")
    assert pkg.pronoun_flag("CL-1", "keep_original", None, {}) is None
    assert pkg.pronoun_flag("CL-1", "propose_change", "A person runs.", {}) is None


def test_the_package_is_a_prior_the_candidate_builder_accepts(candidate, tmp_path):
    source, manifest, path, frames, _, r25 = candidate
    package = Path(_build(candidate, tmp_path)["output"])
    _decide(package, row=2, status="Approved", reviewer="rev")
    first = next(iter(_log(package)))
    out = gt_candidate_build.build(source, manifest, package, path.parent / "ai-proposals-v0.2.jsonl", tmp_path / "v3",
                                   "v0.3", frames_dir=frames)
    assert out["problems"] == [] and out["statuses"]["Approved"] == 1
    assert _log(tmp_path / "v3" / "Ground Truth Candidate v0.3.xlsx")[first]["Review status"] == "Approved"


def test_decisions_carry_forward_into_a_new_package_and_nothing_is_overwritten(candidate, tmp_path):
    first = Path(_build(candidate, tmp_path)["output"])
    with pytest.raises(review_backup.OverwriteRefusedError):
        _build(candidate, tmp_path)
    _decide(first, row=2, status="Rejected", reviewer="rev")
    change_id = next(iter(_log(first)))
    second = _build(candidate, tmp_path, version="v0.2.2", prior_package=first)
    assert second["carried_decisions"] == 1
    assert _log(Path(second["output"]))[change_id]["Review status"] == "Rejected"


def test_the_package_holds_no_annotation_text_in_git_tracked_outputs(candidate, tmp_path):
    result = _build(candidate, tmp_path)
    verification = (tmp_path / "html" / "verification-v0.2.1.json").read_text()
    assert "zebra-striped umbrella" not in verification and "annotator" not in verification
    assert Path(result["output"]).suffix == ".xlsx"  # gitignored (ground_truth/review/*.xlsx)
