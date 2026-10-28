# SPDX-FileCopyrightText: Copyright (c) 2025-2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""AI-assisted correction and the candidate workbook (PR 4), on the synthetic PR 3 workbook: no real data, no video."""

from __future__ import annotations

import copy
import hashlib
import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import ground_truth_correction as gtc  # noqa: E402
import ground_truth_review as gtr  # noqa: E402
import gt_candidate_build  # noqa: E402
import gt_review_build  # noqa: E402
import gt_review_packet  # noqa: E402
import review_backup  # noqa: E402
from test_ground_truth_review import _decide, _manifest, _workbook  # noqa: E402

FRAMES = {"sheets": [{"file": "window-1.jpg", "times": [0.0, 1.0, 2.0, 3.0], "sha256": "a" * 64}]}


@pytest.fixture
def setup(tmp_path):
    source, manifest = _workbook(tmp_path / "source.xlsx"), _manifest(tmp_path / "manifest.json")
    gt_review_build.build(source, manifest, tmp_path / "review", "v0.1")
    return source, manifest, tmp_path / "review" / "Ground Truth Review v0.1.xlsx"


def _changes(source, manifest):
    values = openpyxl.load_workbook(source, data_only=True)
    return values, json.loads(manifest.read_text()), gtr.audit(values, json.loads(manifest.read_text()))[0]


def _first(changes, rule):
    return next(c for c in changes if c.rule == rule)


def _proposal(change, **overrides):
    proposal = {"change_id": change.change_id, "incident_id": change.incident_id, "record_id": str(change.record_id),
                "field": change.field, "decision": "propose_change", "proposed_value": "A person hits another person.",
                "explanation": "No visible gender cues at this resolution.",
                "evidence": [{"sheet": "window-1.jpg", "t": 2.0, "observation": "two figures"}], "confidence": "medium",
                "reviewer": {"kind": "ai", "model": "test-model", "method": "frames"},
                "correction_prompt": gtc.CORRECTION_PROMPT_VERSION, "correction_prompt_sha256": "b" * 64,
                "guidelines": gtr.GUIDELINE_VERSION, "frames_sha256": {"window-1.jpg": "a" * 64}}
    proposal.update(overrides)
    return proposal


def _resolve(changes, values, manifest, prior=None, proposals=None, frames=FRAMES):
    proposals = proposals or {}
    by_id = {c.change_id: c for c in changes}
    verification = {c.change_id: gtc.verify_mechanical(c, values, manifest) for c in changes
                    if c.change_type == "verified_correction"}
    problems = {cid: gtc.validate_proposal(p, by_id[cid], frames) for cid, p in proposals.items()}
    return gtc.resolve(changes, prior=prior, prior_name="prior.xlsx" if prior else None, verification=verification,
                       proposals=proposals, proposal_problems=problems)


def test_every_mechanical_correction_is_independently_verified(setup):
    values, manifest, changes = _changes(*setup[:2])
    verified = [c for c in changes if c.change_type == "verified_correction"]
    assert verified and all(gtc.verify_mechanical(c, values, manifest)[0] for c in verified)


@pytest.mark.parametrize("rule", ["R01", "R02", "R03", "R05", "R28"])
def test_a_tampered_mechanical_correction_fails_verification(setup, rule):
    values, manifest, changes = _changes(*setup[:2])
    change = copy.deepcopy(_first(changes, rule))
    if rule == "R02":
        change.row = 2  # nothing earlier to duplicate
    elif rule == "R05":
        change.original_value = "something else"
    else:
        change.proposed_value = "tampered"
    assert gtc.verify_mechanical(change, values, manifest)[0] is False


def test_only_verified_corrections_apply_automatically(setup):
    values, manifest, changes = _changes(*setup[:2])
    r25 = _first(changes, "R25")
    entries = {e.change_id: e for e in _resolve(changes, values, manifest, proposals={r25.change_id: _proposal(r25)})}
    for change in changes:
        entry = entries[change.change_id]
        if change.change_type == "verified_correction":
            assert entry.status == "Auto-verified" and entry.applied
        elif change.change_type == "informational":
            assert entry.status == "Informational" and not entry.applied
        else:
            assert entry.status == "Pending" and not entry.applied  # rule-based and AI proposals wait for a person
    assert entries[r25.change_id].proposed_value == "A person hits another person."


def test_human_decisions_win_and_an_accepted_value_overrides(setup):
    values, manifest, changes = _changes(*setup[:2])
    r01, r25, r23 = _first(changes, "R01"), _first(changes, "R25"), _first(changes, "R23")
    prior = {r01.change_id: {"Review status": "Rejected", "Reviewer": "rev"},
             r25.change_id: {"Review status": "Approved", "Reviewer": "rev"},
             r23.change_id: {"Review status": "Approved", "Reviewer": "rev", "Accepted value": 2}}
    entries = {e.change_id: e for e in _resolve(changes, values, manifest, prior=prior,
                                                proposals={r25.change_id: _proposal(r25)})}
    assert entries[r01.change_id].status == "Rejected" and not entries[r01.change_id].applied
    assert entries[r25.change_id].applied and entries[r25.change_id].applied_value == "A person hits another person."
    assert entries[r23.change_id].applied and entries[r23.change_id].applied_value == 2
    assert entries[r23.change_id].carried_from == "prior.xlsx" and entries[r23.change_id].reviewer == "rev"


def test_approving_a_review_entry_without_a_proposal_confirms_the_original(setup):
    values, manifest, changes = _changes(*setup[:2])
    r25 = _first(changes, "R25")
    keep = _proposal(r25, decision="keep_original", proposed_value=None)
    entries = {e.change_id: e for e in _resolve(changes, values, manifest,
                                                prior={r25.change_id: {"Review status": "Approved"}},
                                                proposals={r25.change_id: keep})}
    assert entries[r25.change_id].status == "Approved" and not entries[r25.change_id].applied


@pytest.mark.parametrize("override, problem", [
    ({"decision": "maybe"}, "decision"),
    ({"confidence": "certain"}, "confidence"),
    ({"evidence": []}, "evidence timestamp"),
    ({"evidence": [{"sheet": "window-1.jpg", "t": 9.0}]}, "not a tile"),
    ({"evidence": [{"sheet": "overview.jpg", "t": 1.0}]}, "not in frames.json"),
    ({"frames_sha256": {"window-1.jpg": "c" * 64}}, "changed since"),
    ({"decision": "uncertain"}, "must not carry"),
    ({"correction_prompt_sha256": ""}, "missing"),
    ({"record_id": "E9"}, "does not match"),
])
def test_invalid_ai_proposals_are_reported_and_never_used(setup, override, problem):
    values, manifest, changes = _changes(*setup[:2])
    r25 = _first(changes, "R25")
    proposal = _proposal(r25, **override)
    assert any(problem in p for p in gtc.validate_proposal(proposal, r25, FRAMES))
    entry = next(e for e in _resolve(changes, values, manifest, prior={r25.change_id: {"Review status": "Approved"}},
                                     proposals={r25.change_id: proposal}) if e.change_id == r25.change_id)
    assert entry.problems and not entry.applied


def test_r22_proposals_set_only_the_out_of_taxonomy_columns(setup):
    values, manifest, changes = _changes(*setup[:2])
    r22 = _first(changes, "R22")
    good = _proposal(r22, proposed_value={"Out_Of_Taxonomy": True, "Nearest_Type": "explosion"})
    assert gtc.validate_proposal(good, r22, FRAMES) == []
    assert gtc.validate_proposal(_proposal(r22, proposed_value={"Type": "Fire"}), r22, FRAMES)
    entries = _resolve(changes, values, manifest, prior={r22.change_id: {"Review status": "Approved"}},
                       proposals={r22.change_id: good})
    records = gtc.apply_to_records(gtr.read_data_rows(values), entries)
    incident = next(r for r in records["gt_incidents"] if r["Incident_ID (P_key)"] == r22.incident_id)
    assert incident["Out_Of_Taxonomy"] is True and incident["Type"] == "Explosion"


def test_corrected_records_are_deterministic_and_leave_the_source_alone(setup):
    values, manifest, changes = _changes(*setup[:2])
    rows = gtr.read_data_rows(values)
    snapshot = json.dumps(rows, default=str, sort_keys=True)
    first = gtc.apply_to_records(rows, _resolve(changes, values, manifest))
    second = gtc.apply_to_records(rows, _resolve(changes, values, manifest))
    assert gtc.records_digest(first) == gtc.records_digest(second)
    assert json.dumps(rows, default=str, sort_keys=True) == snapshot
    incidents = [r["Incident_ID (P_key)"] for r in first["gt_incidents"]]
    assert incidents.count("Assault018") == 1  # the duplicate row is dropped
    assert any(r["Filename"].endswith("_demo.mp4") for r in first["gt_incidents"])  # kept as annotated
    assert any(r["_change_ids"] for r in first["gt_incidents"])


def test_a_reclassification_must_be_decided_as_a_pair(setup):
    values, manifest, changes = _changes(*setup[:2])
    r10 = [c for c in changes if c.rule == "R10"]
    half = {r10[0].change_id: {"Review status": "Approved"}}
    assert gtc.consistency_problems(_resolve(changes, values, manifest, prior=half))
    both = {c.change_id: {"Review status": "Approved"} for c in r10}
    entries = _resolve(changes, values, manifest, prior=both)
    assert gtc.consistency_problems(entries) == []
    records = gtc.apply_to_records(gtr.read_data_rows(values), entries)
    assert "suv" not in [r["Name"] for r in records["gt_instruments"]]
    assert "suv" in [r["Name"] for r in records["gt_assets"]]


@pytest.mark.parametrize("forbidden", ["eval_data/results/run.json", "contracts/experiments/p1/x.txt",
                                       "eval/experiments/configs/c.json"])
def test_correction_inputs_never_include_system_under_test_outputs(tmp_path, forbidden):
    with pytest.raises(gtc.IndependenceError):
        gtc.guard_inputs([tmp_path / forbidden])
    gtc.guard_inputs([tmp_path / "eval_data" / "gt_frames" / "A" / "frames.json"])


def test_packets_hold_the_annotation_and_findings_but_no_annotator(setup):
    values, manifest, changes = _changes(*setup[:2])
    packets = gt_review_packet.build_packets(values, manifest, changes)
    assert set(packets) == {c.incident_id for c in changes
                            if c.change_type in ("rule_based_proposal", "needs_video_review")}
    text = json.dumps(packets)
    assert "annotator A" not in text and "Labelled_By" not in text
    burglary = packets["Burglary004"]
    assert {f["rule"] for f in burglary["findings"]} >= {"R10", "R20", "R21"}
    assert burglary["window"] == {"start": 5.0, "end": 30.0}


def _candidate(setup, tmp_path, proposals=None, version="v0.2", **kwargs):
    source, manifest, prior = setup
    proposals_path = tmp_path / f"ai-proposals-{version}.jsonl"
    proposals_path.write_text("".join(json.dumps(p) + "\n" for p in proposals or []))
    frames = tmp_path / "frames"
    for incident in ("Assault018", "Burglary004", "Animal004", "Explosion011", "Explosion002", "RoadAccidents005"):
        (frames / incident).mkdir(parents=True, exist_ok=True)
        (frames / incident / "frames.json").write_text(json.dumps(FRAMES))
    return gt_candidate_build.build(source, manifest, prior, proposals_path, tmp_path / "candidate", version,
                                    frames_dir=frames, **kwargs)


def test_candidate_keeps_inputs_unchanged_and_applies_only_verified(setup, tmp_path):
    source, manifest, prior = setup
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, prior)}
    _, _, changes = _changes(source, manifest)
    r25 = _first(changes, "R25")
    result = _candidate(setup, tmp_path, proposals=[_proposal(r25)])
    assert {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, prior)} == before
    assert result["statuses"]["Auto-verified"] == sum(c.change_type == "verified_correction" for c in changes)
    assert result["applied"] == result["statuses"]["Auto-verified"] and result["problems"] == []
    wb = openpyxl.load_workbook(tmp_path / "candidate" / "Ground Truth Candidate v0.2.xlsx")
    log = {row[0]: row for row in wb["Change Log"].iter_rows(min_row=2, values_only=True)}
    assert log[r25.change_id][15] == "Pending" and log[r25.change_id][18] == "no"
    assert log[r25.change_id][7] == "A person hits another person."
    incident_row = next(r for r in wb["gt_incidents"].iter_rows(min_row=2, values_only=True) if r[1] == "Assault018")
    assert "man" in incident_row[6]  # the AI rewrite is not applied to the data sheet
    report = (tmp_path / "candidate" / "candidate-v0.2.md").read_text()
    assert "zebra-striped umbrella" not in report and "annotator" not in report


def test_candidate_carries_prior_decisions_forward(setup, tmp_path):
    _decide(setup[2], row=2, status="Rejected", reviewer="reviewer 1")
    result = _candidate(setup, tmp_path)
    wb = openpyxl.load_workbook(tmp_path / "candidate" / "Ground Truth Candidate v0.2.xlsx")
    first = next(wb["Change Log"].iter_rows(min_row=2, values_only=True))
    assert first[15] == "Rejected" and first[16] == "reviewer 1" and first[27] == setup[2].name
    assert result["prior_state"]["decisions"] == 1


def test_candidate_refuses_a_prior_workbook_for_different_findings(setup, tmp_path):
    wb = openpyxl.load_workbook(setup[2])
    wb["Change Log"].cell(2, 6).value = "Something else"
    wb.save(setup[2])
    with pytest.raises(ValueError, match="does not match"):
        _candidate(setup, tmp_path)


def test_candidate_with_human_decisions_is_never_overwritten(setup, tmp_path):
    _candidate(setup, tmp_path)
    with pytest.raises(review_backup.OverwriteRefusedError):
        _candidate(setup, tmp_path)  # exists, no --overwrite
    _candidate(setup, tmp_path, overwrite=True)  # Auto-verified alone is not a human decision
    _decide(tmp_path / "candidate" / "Ground Truth Candidate v0.2.xlsx", row=2, status="Approved")
    with pytest.raises(review_backup.OverwriteRefusedError, match="review decision"):
        _candidate(setup, tmp_path, overwrite=True)


def test_candidate_cannot_overwrite_its_inputs(setup, tmp_path):
    source, manifest, _ = setup
    _candidate(setup, tmp_path)
    candidate = tmp_path / "candidate" / "Ground Truth Candidate v0.2.xlsx"
    with pytest.raises(review_backup.OverwriteRefusedError, match="new file"):
        gt_candidate_build.build(source, manifest, candidate, tmp_path / "p.jsonl", tmp_path / "candidate", "v0.2",
                                 overwrite=True, frames_dir=tmp_path)


def test_filename_is_never_modified_even_when_approved_with_a_value(setup):
    values, manifest, changes = _changes(*setup[:2])
    r01 = _first(changes, "R01")
    assert r01.change_type == "informational" and r01.proposed_value is None
    prior = {r01.change_id: {"Review status": "Approved", "Accepted value": "Renamed_x264.mp4"}}
    entry = next(e for e in _resolve(changes, values, manifest, prior=prior) if e.change_id == r01.change_id)
    assert not entry.applied and entry.applied_value is None
    records = gtc.apply_to_records(gtr.read_data_rows(values), _resolve(changes, values, manifest, prior=prior))
    assert {r["Filename"] for r in records["gt_incidents"]} >= {"Assault018_x264_demo.mp4"}
    forced = gtc.Entry("CL-9", "R03", "verified_correction", "gt_incidents", 2, "Assault018", "Assault018", "Filename",
                       "a", "b", status="Approved", applied=True, applied_value="b")
    with pytest.raises(ValueError, match="protected"):
        gtc.apply_to_records(gtr.read_data_rows(values), [forced])
