# SPDX-FileCopyrightText: Copyright (c) 2025-2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Ground-truth audit and review workbook (PR 3), on a synthetic workbook: no real annotations, no video."""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import ground_truth_review as gtr  # noqa: E402
import gt_review_build  # noqa: E402

INCIDENT_HEADER = ["Filename", "Incident_ID (P_key)", "Type", "Start_Timestamp", "End_Timestamp", "Duration",
                   "Description", "Severity_Level", "Labelled_By", "Labelled_DateTime"]
UNIQUE_TEXT = "zebra-striped umbrella"  # appears only in a description; must never reach the audit report


def _workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    inc = wb.active
    inc.title = "gt_incidents"
    inc.append(INCIDENT_HEADER)
    when = dt.datetime(2026, 9, 23)
    rows = [
        ("Assault018_x264_demo.mp4", "Assault018", "Assault", 2.0, 9.0, 7.0,
         f"A man in a {UNIQUE_TEXT} hits another man.", 3.0, "annotator A", when),
        ("Assault018_x264.mp4", "Assault018", "Assault", 2.0, 9.0, 7.0,
         f"A man in a {UNIQUE_TEXT} hits another man.", 3.0, "annotator A", when),          # duplicate after R01
        ("Animal004_x264.mp4", "Animal004", "Animal", 1.0, 5.0, 4.0, "A leopard walks past.", 1.0, "annotator B", when),
        ("Explosion011_x264.mp4", "Explosion011", "Explosion", 0.0, 8.0, 8.0, "Smoke comes from a bus.", 2.0,
         "annotator A", when),
        ("Burglary004_x264.mp4", "Burglary004", "Burglary", 5.0, 30.0, 20.0, "The owner forces a door.", 3.0,
         "annotator C", when),                                                               # duration wrong; "owner"
        ("RoadAccidents005_x264.mp4", "RoadAccidents005", "Road Accident", 1.0, 99.0, 98.0, "Two cars collide.", 3.0,
         "annotator A", when),                                                               # end beyond the video
        ("Explosion002_x264.mp4", "Explosion002", "Burglary", 1.0, 4.0, 3.0, "A fire ignites at a pump.", 3.0,
         "annotator D", when),                                                               # filename/type mismatch
    ]
    for row in rows:
        inc.append(list(row))
    inc.append([None, None, None, None, None, 0, None, None, None, None])                   # formatting artifact
    dv = DataValidation(type="list", formula1='"Assault,Burglary,Explosion,Fighting,Animal,Road Accident,Shooting"')
    inc.add_data_validation(dv)
    dv.add("C2:C20")

    ent = wb.create_sheet("gt_entities")
    ent.append(["Filename", "Incident_ID (P_key, F_key)", "Entity_ID (P_key)", "Type", "Description", "Image"])
    ent.append(["Assault018_x264.mp4", "Assault018", "E1", "human", "Tall man in a red jacket", None, "stray,cells"])
    ent.append(["Assault018_x264.mp4", "Assault018", "E2", "human", "Person in grey", None])
    ent.append(["Burglary004_x264.mp4", "Burglary004", "E1", "human", "Person in a hoodie", None])

    ins = wb.create_sheet("gt_instruments")
    ins.append(["Filename", "Incident_ID (P_key, F_key)", "Instrument_ID (P_key)", "Entity_ID", "Name", "Description",
                "Threat_Level", "Image"])
    ins.append(["Burglary004_x264.mp4", "Burglary004", "I1", "E1", "suv", "Silver SUV used to ram the gate", 3.0, None])
    ins.append(["Burglary004_x264.mp4", "Burglary004", "I2", "E1 ", "explosives", "Package", 3.0, None])
    ins.append(["Burglary004_x264.mp4", "Burglary004", "I3", "E1", "red item", "Small red thing", 1.0, None])
    ins.append(["Explosion002_x264.mp4", "Explosion002", "I1", "E7", "fire extinguisher", "Red extinguisher", 1.0, None])

    ast = wb.create_sheet("gt_assets")
    ast.append(["Filename", "Incident_ID (P_key, F_key)", "Asset_ID (P_key)", "Name", "Description", "Image"])
    ast.append(["Burglary004_x264.mp4", "Burglary004", "A1", "door", "Shop door", None])
    ast.append(["Burglary004_x264.mp4", "Burglary004", "A3", "atm", "ATM", None])                # id gap
    wb.create_sheet("Severity Guide").append(["Category", "Level"])
    hidden = wb.create_sheet("Metadata (hidden)")
    hidden.sheet_state = "hidden"
    wb.save(path)
    return path


def _manifest(path: Path) -> Path:
    clips = []
    for incident_id, folder, demo, duration in (("Assault018", "demo", True, 20.0), ("Animal004", "animal_attacks", False, 9.0),
                                                ("Explosion011", "explosion", False, 12.0),
                                                ("Burglary004", "burglary", False, 40.0),
                                                ("RoadAccidents005", "road_accidents", False, 30.0),
                                                ("Explosion002", "explosion", False, 10.0)):
        basename = f"{incident_id}_x264{'_demo' if demo else ''}.mp4"
        clips.append({"incident_id": incident_id, "filename": f"{incident_id}_x264.mp4",
                      "r2_key": f"anomaly/{folder}/{basename}", "local_path": f"eval_data/videos/{folder}/{basename}",
                      "duration_seconds": duration})
    path.write_text(json.dumps({"clips": clips, "clips_sha256": "0" * 64}))
    return path


@pytest.fixture
def sources(tmp_path):
    return _workbook(tmp_path / "source.xlsx"), _manifest(tmp_path / "manifest.json")


def _audit(sources):
    workbook, manifest = sources
    return gtr.audit(openpyxl.load_workbook(workbook, data_only=True), json.loads(manifest.read_text()))


def _by_rule(changes):
    out = {}
    for change in changes:
        out.setdefault(change.rule, []).append(change)
    return out


def test_every_rule_fires_with_its_class(sources):
    rules = _by_rule(_audit(sources)[0])
    assert {"R01", "R02", "R03", "R05", "R10", "R20", "R21", "R22", "R23", "R24", "R25", "R26", "R28", "R29", "R30",
            "R34"} <= set(rules)
    for rule, changes in rules.items():
        assert {c.change_type for c in changes} == {gtr.RULES[rule][0]}
    assert rules["R01"][0].proposed_value == "Assault018_x264.mp4"
    assert rules["R03"][0].proposed_value == "E1"
    assert rules["R02"][0].proposed_value == "remove (duplicate of row 2)"


def test_no_value_is_proposed_where_only_the_video_can_answer(sources):
    for change in _audit(sources)[0]:
        if change.change_type in ("needs_video_review", "informational"):
            assert change.proposed_value is None
            assert change.change_type == "informational" or change.evidence.startswith("Required: watch")


def test_types_are_never_changed_and_out_of_taxonomy_is_a_flag(sources):
    changes = _audit(sources)[0]
    assert not [c for c in changes if c.field == "Type" and c.proposed_value is not None]
    oot = [c for c in changes if c.rule == "R22"]
    assert [c.incident_id for c in oot] == ["Explosion011"] and oot[0].review_flag == "OUT_OF_TAXONOMY_CANDIDATE"
    # Explosion002 is labelled Burglary: the filename prefix is not evidence, so nothing is proposed for it.
    assert not [c for c in changes if c.incident_id == "Explosion002" and c.field == "Type"]


def test_rows_proposed_for_removal_get_no_other_edits(sources):
    changes = _audit(sources)[0]
    duplicate_rows = {c.row for c in changes if c.rule == "R02"}
    assert duplicate_rows == {3}
    assert [c.rule for c in changes if c.sheet == "gt_incidents" and c.row in duplicate_rows] == ["R02"]


def test_vehicle_instrument_becomes_a_new_asset_with_the_next_free_id(sources):
    rules = _by_rule(_audit(sources)[0])
    removal, addition = rules["R10"]
    assert removal.review_flag == "PROPOSED_REMOVAL_RECLASSIFIED" and addition.record_id == "A4"
    assert addition.proposed_value["Name"] == "suv" and "Threat_Level" not in addition.proposed_value


def test_build_keeps_the_source_and_its_records(sources, tmp_path):
    workbook, manifest = sources
    before = hashlib.sha256(workbook.read_bytes()).hexdigest()
    result = gt_review_build.build(workbook, manifest, tmp_path / "out", "v0.1")
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == before == result["source_sha256"]

    src = openpyxl.load_workbook(workbook)
    out = openpyxl.load_workbook(tmp_path / "out" / "Ground Truth Review v0.1.xlsx")
    assert out.sheetnames[:2] == ["Review Summary", "Change Log"]
    assert out.sheetnames[2:] == src.sheetnames and out["Metadata (hidden)"].sheet_state == "hidden"
    for sheet, spec in gtr.SHEETS.items():
        documented = len(spec["columns"])
        assert [c.value for c in out[sheet][1]][:documented] == spec["columns"]
        assert [c.value for c in out[sheet][1]][documented:documented + 2] == ["review_flags", "change_ids"]
        id_col = spec["columns"].index(spec["id"]) + 1
        src_ids = [src[sheet].cell(r, id_col).value for r in range(2, src[sheet].max_row + 1)]
        out_ids = [out[sheet].cell(r, id_col).value for r in range(2, len(src_ids) + 2)]
        assert out_ids == src_ids  # same rows, same order: nothing deleted or renumbered
    assert len(out["gt_incidents"].data_validations.dataValidation) == 1
    assert out["gt_assets"].cell(4, 3).value == "A4" and out["gt_assets"].cell(4, 7).value == "NEW_RECORD"
    assert out["gt_incidents"].cell(3, 11).value == "PROPOSED_REMOVAL_DUPLICATE"
    assert out["gt_incidents"].cell(2, 1).value == "Assault018_x264.mp4"
    assert out["gt_instruments"].cell(3, 4).value == "E1"
    assert out["gt_entities"].cell(2, 7).value is None or out["gt_entities"].cell(2, 7).value.startswith(("NEEDS", ""))
    assert [c.value for c in out["gt_incidents"]["C"][1:8]] == [c.value for c in src["gt_incidents"]["C"][1:8]]


def test_change_log_and_summary(sources, tmp_path):
    workbook, manifest = sources
    result = gt_review_build.build(workbook, manifest, tmp_path / "out", "v0.1")
    out = openpyxl.load_workbook(tmp_path / "out" / "Ground Truth Review v0.1.xlsx")
    log = out["Change Log"]
    assert [c.value for c in log[1]] == gt_review_build.LOG_COLUMNS
    assert log.max_row - 1 == result["entries"]
    assert {log.cell(r, 16).value for r in range(2, log.max_row + 1)} == {"Pending"}
    validation = log.data_validations.dataValidation[0]
    assert validation.formula1 == '"Pending,Approved,Rejected"' and str(validation.sqref) == f"P2:P{log.max_row}"
    summary = out["Review Summary"]
    formulas = [c.value for row in summary.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("'Change Log'!$P$2" in f and '"Approved"' in f for f in formulas)
    assert any('"needs_video_review"' in f for f in formulas)


def test_build_is_reproducible_and_the_report_holds_no_annotation_text(sources, tmp_path):
    workbook, manifest = sources
    first = gt_review_build.build(workbook, manifest, tmp_path / "one", "v0.1")
    second = gt_review_build.build(workbook, manifest, tmp_path / "two", "v0.1")
    assert first["change_log_sha256"] == second["change_log_sha256"]
    report = (tmp_path / "one" / "audit-v0.1.md").read_text()
    assert UNIQUE_TEXT not in report
    assert not any(name in report for name in ("annotator A", "annotator B", "annotator C", "annotator D"))
    assert first["source_sha256"] in report
    manifest_out = json.loads((tmp_path / "one" / "review-v0.1-manifest.json").read_text())
    assert manifest_out["entries"] == first["entries"] and manifest_out["output_sha256"]


def test_appended_rows_carry_formula_columns(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "gt_assets"
    ws.append(["Filename", "Incident_ID (P_key, F_key)", "Asset_ID (P_key)", "Name", "Description", "Image"])
    ws.append(["Burglary004_x264.mp4", '=LEFT(A2,FIND("_",A2) - 1)', "A1", "door", "Door", None])
    for sheet in ("gt_incidents", "gt_entities", "gt_instruments"):
        wb.create_sheet(sheet).append(gtr.SHEETS[sheet]["columns"])
    change = gtr.Change("R10", "gt_assets", 0, "Burglary004", "A2", "(new row)", None,
                        {"Filename": "Burglary004_x264.mp4", "Incident_ID (P_key, F_key)": "Burglary004",
                         "Asset_ID (P_key)": "A2", "Name": "suv", "Description": "SUV", "Image": None},
                        "r", "e", "High", change_id="CL-0001")
    gt_review_build.apply_changes(wb, [change])
    assert ws["B3"].value == '=LEFT(A3,FIND("_",A3) - 1)' and ws["C3"].value == "A2" and change.row == 3


def test_demo_filename_without_a_manifest_clip_is_not_rewritten(sources):
    workbook, manifest = sources
    data = json.loads(manifest.read_text())
    data["clips"] = [c for c in data["clips"] if c["incident_id"] != "Assault018"]
    changes, facts = gtr.audit(openpyxl.load_workbook(workbook, data_only=True), data)
    assert not [c for c in changes if c.rule == "R01"]
    assert "Assault018" in facts["incidents_without_clip"]
