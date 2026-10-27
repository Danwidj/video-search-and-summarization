# SPDX-FileCopyrightText: Copyright (c) 2025-2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Offline (database-free) scoring: the same scores as the database path, and no database access.

The equivalence tests write one synthetic workbook and the same predictions into the
SQLite ``IncidentDB`` the existing scorer tests use and into an ``OfflineEvalStore``,
score both with the unchanged ``eval_gt.run_evaluation``, and require identical results.
The guard tests make every database entry point raise and show the default path never
reaches one, even with finals-style credentials in the environment.
"""

from __future__ import annotations

import dataclasses
import datetime as dt
import re
import sys
from collections import Counter
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import db as db_module  # noqa: E402
import db_postgrest  # noqa: E402
import embed_client  # noqa: E402
import eval_gt  # noqa: E402
import eval_ingest_gt  # noqa: E402
import eval_run  # noqa: E402
import matching  # noqa: E402
from eval_gt import Result, run_evaluation  # noqa: E402
from offline_store import OfflineEvalStore  # noqa: E402

_INCIDENT_HEADER = ["Filename", "Incident_ID (P_key)", "Type", "Start_Timestamp", "End_Timestamp", "Duration",
                    "Description", "Severity_Level", "Labelled_By", "Labelled_DateTime"]
_ENTITY_HEADER = ["Filename", "Incident_ID (P_key, F_key)", "Entity_ID (P_key)", "Type", "Description", "Image"]
_INSTRUMENT_HEADER = ["Filename", "Incident_ID (P_key, F_key)", "Instrument_ID (P_key)", "Entity_ID", "Name",
                      "Description", "Threat_Level", "Image"]
_ASSET_HEADER = ["Filename", "Incident_ID (P_key, F_key)", "Asset_ID (P_key)", "Name", "Description", "Image"]

_LABELLED = dt.datetime(2026, 9, 23, 10, 30)

# Float cells and capitalised / retired types exactly as the real workbook stores them.
_GT_INCIDENTS = [
    ("Burglary001_x264.mp4", "Burglary001", "Burglary", 12.0, 40.0, 28.0,
     "A man in a red jacket forces a shop door with a crowbar and takes the cash box.", 3.0),
    ("Animal002_x264.mp4", "Animal002", "Animal", 3.0, 9.0, 6.0,
     "A large dog attacks a cyclist on a footpath.", 4.0),
    ("Explosion003_x264.mp4", "Explosion003", "Explosion", 0.0, 15.0, 15.0,
     "A gas cylinder explodes beside a parked van.", 5.0),
]
_GT_ENTITIES = [
    ("Burglary001", "E1", "human", "man in a red jacket forcing the door"),
    ("Burglary001", "E2", "human", "shop owner arriving later"),
    ("Burglary001", "E10", "human", "passer by who calls for help"),
    ("Animal002", "E1", "animal", "large brown dog"),
    ("Animal002", "E2", "human", "cyclist knocked off the bicycle"),
]
_GT_INSTRUMENTS = [
    ("Burglary001", "I1", "E1", "crowbar", "metal crowbar used to force the door", 3.0),
    ("Burglary001", "I2", "E1 ", "bag", "bag holding the cash box", 1.0),  # trailing space, as in the workbook
]
_GT_ASSETS = [
    ("Burglary001", "A1", "door", "shop door forced open"),
    ("Burglary001", "A2", "cash box", "cash box taken from the counter"),
    ("Animal002", "A1", "bicycle", "bicycle knocked over"),
    ("Explosion003", "A1", "van", "parked white van damaged by the blast"),
]


def _write_workbook(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    incidents = workbook.active
    incidents.title = "gt_incidents"
    incidents.append(_INCIDENT_HEADER)
    for filename, incident_id, kind, start, end, duration, description, severity in _GT_INCIDENTS:
        incidents.append([filename, incident_id, kind, start, end, duration, description, severity, "annotator A",
                          _LABELLED])
    incidents.append([None, None, None, None, None, 0, None, None, None, None])  # formatting-artifact row
    entities = workbook.create_sheet("gt_entities")
    entities.append(_ENTITY_HEADER)
    for incident_id, entity_id, kind, description in _GT_ENTITIES:
        entities.append([f"{incident_id}_x264.mp4", incident_id, entity_id, kind, description, None])
    instruments = workbook.create_sheet("gt_instruments")
    instruments.append(_INSTRUMENT_HEADER)
    for incident_id, instrument_id, holder, name, description, threat in _GT_INSTRUMENTS:
        instruments.append([f"{incident_id}_x264.mp4", incident_id, instrument_id, holder, name, description, threat,
                            None])
    assets = workbook.create_sheet("gt_assets")
    assets.append(_ASSET_HEADER)
    for incident_id, asset_id, name, description in _GT_ASSETS:
        assets.append([f"{incident_id}_x264.mp4", incident_id, asset_id, name, description, None])
    workbook.save(path)
    return path


def _prediction(incident: dict, entities=(), instruments=(), assets=()) -> dict:
    return {"incident": incident, "entities": list(entities), "instruments": list(instruments), "assets": list(assets)}


# (incident id, prediction) pairs covering every scoring branch: exact and near matches, a
# false positive, false negatives, an unresolved holder, ids that sort differently as text
# (E10 before E2), float severities, and an empty prediction (an invalid output scored as a miss).
_PREDICTIONS = [
    ("Burglary001", _prediction(
        {"type": "burglary", "start_timestamp": 14, "end_timestamp": 39, "duration": 25,
         "description": "A man in a red jacket breaks a shop door with a crowbar and steals a cash box.",
         "severity_level": 3, "confidence_score": None},
        entities=[{"entity_id": "E1", "type": "human", "description": "man in a red jacket forcing the door"},
                  {"entity_id": "E2", "type": "human", "description": "dog walking past"},
                  {"entity_id": "E10", "type": "human", "description": "passer by who calls for help"}],
        instruments=[{"entity_id": "E1", "instrument_id": "I1", "name": "crowbar",
                      "description": "metal crowbar used to force the door", "threat_level": 4.0},
                     {"entity_id": "E9", "instrument_id": "I2", "name": "torch", "description": "small torch",
                      "threat_level": 1}],
        assets=[{"asset_id": "A1", "name": "door", "description": "shop door forced open"}],
    )),
    ("Animal002", _prediction(
        {"type": "animal attack", "start_timestamp": 20, "end_timestamp": 30, "duration": 10,
         "description": "A dog bites a cyclist.", "severity_level": 4.0, "confidence_score": 0.7},
        entities=[{"entity_id": "E1", "type": "animal", "description": "large brown dog"}],
    )),
    ("Explosion003", eval_run.EMPTY_PREDICTION),
]


def _fake_embed(texts):
    """Deterministic bag-of-words term-count vectors, no network (as in test_eval_gt.py)."""
    tokenized = [re.findall(r"[a-z']+", text.lower()) for text in texts]
    vocab = sorted({word for tokens in tokenized for word in tokens})
    return embed_client.Result(ok=True, data=[[float(Counter(t).get(w, 0)) for w in vocab] for t in tokenized])


def _fake_judge(expected, predicted, **_):
    """Deterministic token-overlap score, so every judged pair gets its own value."""
    a, b = set(re.findall(r"[a-z]+", (expected or "").lower())), set(re.findall(r"[a-z]+", (predicted or "").lower()))
    return Result(ok=True, data=round(len(a & b) / len(a | b), 6) if a | b else 0.0)


def _load_and_score(store, workbook_path: Path, *, incident_only: bool) -> dict[str, dict]:
    eval_ingest_gt.write_gt(store, openpyxl.load_workbook(workbook_path, data_only=True))
    store.insert_model_run("MR-1", model_name="m", prompt_version="incident-contract-v4")
    scores = {}
    with (
        patch.object(matching, "embed_texts", side_effect=_fake_embed),
        patch.object(eval_gt, "judge_description_similarity", side_effect=_fake_judge),
    ):
        for incident_id, prediction in _PREDICTIONS:
            eval_run.persist_prediction(store, incident_id, "MR-1", prediction)
            scores[incident_id] = dataclasses.asdict(
                run_evaluation(store, incident_id, "MR-1", incident_only=incident_only)
            )
    return scores


@pytest.mark.parametrize("incident_only", [False, True])
def test_offline_scores_equal_the_database_scorer(incident_db, tmp_path, incident_only):
    workbook_path = _write_workbook(tmp_path / "gt.xlsx")
    database_scores = _load_and_score(incident_db, workbook_path, incident_only=incident_only)
    offline_scores = _load_and_score(OfflineEvalStore(), workbook_path, incident_only=incident_only)
    assert offline_scores == database_scores


def test_equivalence_fixture_exercises_every_scoring_branch(tmp_path):
    scores = _load_and_score(OfflineEvalStore(), _write_workbook(tmp_path / "gt.xlsx"), incident_only=False)
    burglary = scores["Burglary001"]
    assert burglary["fields"]["start_timestamp"]["pass"] is True and burglary["fields"]["type"]["pass"] is True
    assert burglary["counts"]["entities"]["tp"] == 2 and burglary["counts"]["entities"]["fp"] == 1
    assert burglary["counts"]["instruments"]["fp"] == 1 and burglary["counts"]["assets"]["fn"] == 1
    instrument = burglary["matches"]["instruments"][0]["attribute_scores"]
    assert instrument["threat_level"]["pass"] is False and instrument["holder"]["status"] == "correct"
    assert scores["Animal002"]["fields"]["severity_level"]["pass"] is True  # 4.0 == 4
    assert scores["Animal002"]["fields"]["start_timestamp"]["pass"] is False
    explosion = scores["Explosion003"]
    assert explosion["description_score"] == 0.0 and explosion["counts"]["assets"]["fn"] == 1


def test_offline_rows_match_database_rows(incident_db, tmp_path):
    workbook_path = _write_workbook(tmp_path / "gt.xlsx")
    offline = OfflineEvalStore()
    for store in (incident_db, offline):
        eval_ingest_gt.write_gt(store, openpyxl.load_workbook(workbook_path, data_only=True))
        store.insert_model_run("MR-1", model_name="m")
        eval_run.persist_prediction(store, "Burglary001", "MR-1", _PREDICTIONS[0][1])
    assert offline.get_gt_incident("Burglary001") == incident_db.get_gt_incident("Burglary001")
    assert offline.get_incident("Burglary001", "MR-1") == incident_db.get_incident("Burglary001", "MR-1")
    for kind in ("entities", "instruments", "assets"):
        assert getattr(offline, f"list_gt_{kind}")("Burglary001") == getattr(incident_db, f"list_gt_{kind}")(
            "Burglary001"
        )
        assert getattr(offline, f"list_incident_{kind}")("Burglary001", "MR-1") == getattr(
            incident_db, f"list_incident_{kind}"
        )("Burglary001", "MR-1")
    assert [e["entity_id"] for e in offline.list_gt_entities("Burglary001")] == ["E1", "E10", "E2"]
    assert isinstance(offline.get_gt_incident("Burglary001")["start_timestamp"], str)  # VARCHAR column, as in SQL


def test_reinserting_an_incident_replaces_its_evidence_and_matches(tmp_path):
    store = OfflineEvalStore()
    eval_ingest_gt.write_gt(store, openpyxl.load_workbook(_write_workbook(tmp_path / "gt.xlsx"), data_only=True))
    with patch.object(matching, "embed_texts", side_effect=_fake_embed):
        eval_run.persist_prediction(store, "Burglary001", "MR-1", _PREDICTIONS[0][1])
        matching.match_incident(store, "Burglary001", "MR-1")
    assert store.list_matches("entities", "Burglary001", "MR-1")
    eval_run.persist_prediction(store, "Burglary001", "MR-1", eval_run.EMPTY_PREDICTION)
    assert store.list_incident_entities("Burglary001", "MR-1") == []
    assert store.list_matches("entities", "Burglary001", "MR-1") == []
    eval_ingest_gt.write_gt(store, openpyxl.load_workbook(tmp_path / "gt.xlsx", data_only=True))
    assert len(store.list_gt_entities("Burglary001")) == 3  # replaced, not duplicated


# --- the default path never reaches a database ----------------------------------------


@pytest.fixture
def finals_style_env(monkeypatch):
    """Credentials shaped like the finals project's, so any database access would have somewhere to go."""
    monkeypatch.setenv("INCIDENT_DB_DSN", "postgresql://postgres.finals:secret@aws-0-ap-southeast-1.pooler.supabase.com:5432/postgres")
    monkeypatch.setenv("INCIDENT_SUPABASE_URL", "https://finals-project.supabase.co")
    monkeypatch.setenv("INCIDENT_SUPABASE_SERVICE_ROLE_KEY", "placeholder-service-role-key")
    db_module.reset_cache()


@pytest.fixture
def database_forbidden(monkeypatch):
    """Every way eval reaches a database raises; ``calls`` lists any attempt."""
    calls: list[str] = []

    def forbid(name):
        def _raise(*args, **kwargs):
            calls.append(name)
            raise AssertionError(f"database access attempted: {name}")
        return _raise

    monkeypatch.setattr(db_module, "get_db", forbid("db.get_db"))
    monkeypatch.setattr(db_module.IncidentDB, "from_dsn", forbid("IncidentDB.from_dsn"))
    monkeypatch.setattr(db_module.IncidentDB, "__init__", forbid("IncidentDB()"))
    monkeypatch.setattr(db_module.IncidentDB, "init_schema", forbid("IncidentDB.init_schema"))
    monkeypatch.setattr(db_module, "build_engines", forbid("db.build_engines"))
    monkeypatch.setattr(db_postgrest.PostgrestIncidentDB, "__init__", forbid("PostgrestIncidentDB()"))
    monkeypatch.setattr(eval_run, "PostgrestIncidentDB", forbid("eval_run.PostgrestIncidentDB"))
    monkeypatch.setattr("sqlalchemy.create_engine", forbid("sqlalchemy.create_engine"))
    monkeypatch.setattr("httpx.Client.__init__", forbid("httpx.Client() (PostgREST)"))
    return calls


def test_default_scoring_store_is_offline_and_touches_no_database(
    finals_style_env, database_forbidden, tmp_path, monkeypatch
):
    monkeypatch.setattr(eval_ingest_gt, "WORKBOOK_PATH", _write_workbook(tmp_path / "gt.xlsx"))
    store = eval_run.scoring_store()
    assert isinstance(store, OfflineEvalStore)
    assert store.get_gt_incident("Burglary001")["type"] == "burglary"
    assert database_forbidden == []


def test_offline_run_writes_and_scores_without_any_database(finals_style_env, database_forbidden, tmp_path,
                                                            monkeypatch):
    monkeypatch.setattr(eval_ingest_gt, "WORKBOOK_PATH", _write_workbook(tmp_path / "gt.xlsx"))
    store = eval_run.scoring_store()
    store.insert_model_run("MR-1", model_name="m")
    with (
        patch.object(matching, "embed_texts", side_effect=_fake_embed),
        patch.object(eval_gt, "judge_description_similarity", side_effect=_fake_judge),
    ):
        for incident_id, prediction in _PREDICTIONS:
            eval_run.persist_prediction(store, incident_id, "MR-1", prediction)
            run_evaluation(store, incident_id, "MR-1")
    assert store.list_matches("entities", "Burglary001", "MR-1")
    assert database_forbidden == []


def test_run_defaults_to_the_offline_store(finals_style_env, database_forbidden, tmp_path, monkeypatch):
    """``eval_run.run`` end to end (P1, R2 and the judge faked): offline backend, no database access."""
    import json

    monkeypatch.setattr(eval_ingest_gt, "WORKBOOK_PATH", _write_workbook(tmp_path / "gt.xlsx"))
    (tmp_path / "split_burglary.json").write_text(json.dumps({
        "category": "Burglary", "evaluation_videos": ["Burglary001_x264.mp4"],
        "evaluation_videos_r2_keys": {"Burglary001_x264.mp4": "anomaly/burglary/Burglary001_x264.mp4"},
    }))
    gt_index = {"by_category": {"Burglary": []},
                "by_incident_id": {"Burglary001": {"filename": "Burglary001_x264.mp4", "p1_shaped": {}}}}
    written = []

    def fake_run_one_video(db, *, incident_id, model_run_id, **kwargs):
        assert isinstance(db, OfflineEvalStore)
        eval_run.persist_prediction(db, incident_id, model_run_id, _PREDICTIONS[0][1])
        return {"filename": kwargs["filename"], "timing": {"p1_elapsed_s": 1.0}, "outcome": "scored",
                "judge_calls": 1, "p1_raw": {"ok": True, "contract_ok": True, "core_ok": True,
                                             "finish_reason": "stop", "rate_limit": {"attempt_count": 1}},
                "incident_field_scores": {}, "rp1_report": {"ok": False}}

    with (
        patch("eval_run.EVAL_DATA_DIR", tmp_path),
        patch("eval_run.model_run_id_for", side_effect=lambda model: f"P1-{model}"),
        patch("eval_gt.judge_description_similarity", return_value=Result(ok=True, data=1.0)),
        patch("eval_run.load_gt_index", return_value=gt_index),
        patch("eval_run.r2_videos.list_video_keys", return_value=["anomaly/burglary/Burglary001_x264.mp4"]),
        patch("eval_run.r2_video_metadata", return_value={"size_bytes": 10, "duration_seconds": 50.0}),
        patch("eval_run.signed_url", side_effect=lambda key: f"https://r2.example/{key}?sig"),
        patch("eval_run.run_one_video", side_effect=fake_run_one_video),
        patch("eval_run.write_results", side_effect=lambda payload, **k: (written.append(payload), ("o", "a"))[1]),
    ):
        eval_run.run(models=["m1"], categories=["Burglary"], limit=None, incident_only=True)

    assert written and written[0]["benchmark"]["scoring_backend"] == "offline"
    assert database_forbidden == []


def test_init_schema_needs_the_database_path():
    with pytest.raises(ValueError, match="init_schema requires use_db"):
        eval_run.scoring_store(init_schema=True)


def test_get_db_never_creates_tables_unless_asked(incident_db, monkeypatch):
    created = []
    monkeypatch.setattr("config.incident_db_dsn", lambda: str(incident_db.engine.url))
    monkeypatch.setattr(db_module.IncidentDB, "init_schema", lambda self: created.append(True))
    db_module.reset_cache()
    assert db_module.get_db() is not None
    assert created == []
    db_module.reset_cache()
    db_module.get_db(init_schema=True)
    assert created == [True]


def test_get_db_returns_none_for_an_unreachable_database(monkeypatch, tmp_path):
    monkeypatch.setattr("config.incident_db_dsn", lambda: f"sqlite:///{tmp_path / 'missing' / 'no.db'}")
    db_module.reset_cache()
    assert db_module.get_db() is None


def test_ingest_dry_run_opens_no_database(finals_style_env, database_forbidden, tmp_path, monkeypatch):
    monkeypatch.setattr(eval_ingest_gt, "WORKBOOK_PATH", _write_workbook(tmp_path / "gt.xlsx"))
    monkeypatch.setattr(eval_ingest_gt.r2_videos, "list_video_keys", lambda: [
        f"anomaly/x/{incident_id}_x264.mp4" for _, incident_id, *_ in _GT_INCIDENTS
    ])
    summary = eval_ingest_gt.ingest(dry_run=True)
    assert summary["incidents"] == 3 and summary["entities"] == 5 and summary["unresolved_videos"] == []
    assert database_forbidden == []


def test_eval_scripts_require_an_explicit_flag_for_the_database():
    import subprocess

    scripts = Path(__file__).resolve().parents[1] / "scripts"
    for script, args in (("eval_ingest_gt.py", []), ("eval_run.py", ["--init-schema"]),
                         ("eval_ingest_gt.py", ["--dry-run", "--init-schema"])):
        completed = subprocess.run([sys.executable, str(scripts / script), *args], check=False, capture_output=True, text=True,
                                   timeout=60, env={"PATH": "/usr/bin:/bin"})
        assert completed.returncode == 2, (script, args, completed.stderr)


def test_offline_coerces_values_like_the_database(incident_db):
    """Floats in integer columns, numbers in text columns, ints in float columns: read back as SQL returns them."""
    offline = OfflineEvalStore()
    for store in (incident_db, offline):
        store.upsert_video("Burglary009", filepath="anomaly/burglary/Burglary009_x264.mp4")
        store.insert_model_run("MR-1", model_name="m")
        store.insert_gt_incident("Burglary009", fields={
            "type": "burglary", "start_timestamp": 12.0, "end_timestamp": 40.5, "duration": 28.0,
            "description": "d", "severity_level": 3.0, "labelled_by": "annotator A", "labelled_datetime": _LABELLED,
        })
        store.add_gt_instrument("Burglary009", instrument_id="I1", entity_id="E4 ", name="stick", threat_level=3.0)
        store.insert_incident("Burglary009", "MR-1", fields={
            "type": "burglary", "start_timestamp": 7, "end_timestamp": "00:00:31", "duration": 24,
            "severity_level": 2, "confidence_score": 1,
        })
    assert offline.get_gt_incident("Burglary009") == incident_db.get_gt_incident("Burglary009")
    assert offline.list_gt_instruments("Burglary009") == incident_db.list_gt_instruments("Burglary009")
    assert offline.get_incident("Burglary009", "MR-1") == incident_db.get_incident("Burglary009", "MR-1")
    assert offline.get_gt_incident("Burglary009")["start_timestamp"] == "12.0"
    assert offline.get_gt_incident("Burglary009")["severity_level"] == 3
    assert offline.get_incident("Burglary009", "MR-1")["confidence_score"] == 1.0
