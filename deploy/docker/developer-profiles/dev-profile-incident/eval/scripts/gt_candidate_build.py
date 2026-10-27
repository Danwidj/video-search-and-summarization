# SPDX-FileCopyrightText: Copyright (c) 2025-2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Build a versioned ground-truth candidate workbook from the source, a reviewed workbook and AI proposals (PR 4).

Inputs, none of which is ever written:
- the source workbook (``Complete Ground Truth Labelling.xlsx``), re-audited so every finding
  is recomputed rather than trusted;
- the previous review workbook (``Ground Truth Review v0.1.xlsx``): its decisions (status,
  reviewer, notes, accepted value) are carried forward unchanged, and the build refuses if it
  does not describe the same findings;
- the AI video-review proposals (``ai-proposals-<version>.jsonl``), each validated against its
  finding and the contact sheets it cites.

What the candidate applies to its data sheets:
- ``verified_correction`` entries that pass the independent re-check (``Auto-verified``);
- entries a person marked ``Approved`` (with their accepted value when given).

Nothing else. AI proposals and rule-based proposals are listed with their evidence and
confidence but stay ``Pending`` and unapplied until a person approves them. A candidate is
**not** ground truth: PR 5 freezes only after the final annotations are approved.

Usage (from ``eval/``)::

    uv run python scripts/gt_candidate_build.py --version v0.2
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import ground_truth_correction as gtc  # noqa: E402
import ground_truth_review as gtr  # noqa: E402
import openpyxl  # noqa: E402
import review_backup  # noqa: E402
from gt_review_build import (  # noqa: E402
    DEFAULT_MANIFEST,
    DEFAULT_OUT_DIR,
    DEFAULT_SOURCE,
    EVAL_DIR,
    FILLS,
    HEADER_FILL,
    HEADER_FONT,
    LOG_COLUMNS,
    WRAP,
    _column_index,
    _display,
    _last_data_row,
    sha256_file,
)
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.formula.translate import Translator  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

GENERATOR = "gt_candidate_build.py v1"
DEFAULT_PRIOR = DEFAULT_OUT_DIR / "Ground Truth Review v0.1.xlsx"
FRAMES_DIR = EVAL_DIR / "eval_data" / "gt_frames"
PROMPT_PATH = EVAL_DIR / "ground_truth" / "correction_prompts" / f"{gtc.CORRECTION_PROMPT_VERSION}.md"
EXTRA_COLUMNS = ["Applied in candidate", "Applied value", "Independent verification", "AI decision", "AI confidence",
                 "AI explanation", "AI reviewer and prompt", "AI proposal problems", "Accepted value", "Carried from"]
PENDING_FILL = PatternFill("solid", fgColor="F8CBAD")


def _evidence_text(proposal: dict) -> str:
    return "; ".join(f"{e.get('sheet')} @{e.get('t')}s: {e.get('observation', '')}" for e in proposal.get("evidence") or [])


def _write_value(ws, entry: gtc.Entry) -> None:
    row = entry.row
    if entry.rule == "R05":
        ws.cell(row, int(entry.field.split()[1])).value = None
        return
    if entry.field == "(new row)":
        documented = len(gtr.SHEETS[entry.sheet]["columns"])
        previous = _last_data_row(ws, documented)
        row = previous + 1
        for header, value in entry.applied_value.items():
            column = _column_index(ws, header)
            above = ws.cell(previous, column).value
            if isinstance(above, str) and above.startswith("="):
                origin = f"{get_column_letter(column)}{previous}"
                value = Translator(above, origin=origin).translate_formula(f"{get_column_letter(column)}{row}")
            ws.cell(row, column).value = value
            ws.cell(row, column).fill = FILLS[entry.change_type]
        entry.row = row
        return
    if entry.rule == "R02" or entry.field == "(row)":
        for column in range(1, len(gtr.SHEETS[entry.sheet]["columns"]) + 1):
            ws.cell(row, column).font = Font(strike=True, color="7F7F7F")
        return
    if entry.rule == "R22":
        for key, value in entry.applied_value.items():
            column = _ensure_column(ws, key)
            ws.cell(row, column).value = value
        return
    cell = ws.cell(row, _column_index(ws, entry.field))
    cell.value = entry.applied_value
    cell.fill = FILLS[entry.change_type]
    cell.comment = Comment(f"{entry.change_id} ({entry.rule}, {entry.status}): was {_display(entry.original_value)!r}",
                           GENERATOR)


def _ensure_column(ws, header: str) -> int:
    """The column titled ``header``, else the first wholly empty column right of the documented ones."""
    try:
        return _column_index(ws, header)
    except KeyError:
        column = len(gtr.SHEETS[ws.title]["columns"]) + 1
        while any(ws.cell(r, column).value not in (None, "") for r in range(1, ws.max_row + 1)):
            column += 1  # never write over a cell, e.g. a stray value whose removal was rejected
        ws.cell(1, column, header).font = Font(bold=True)
        return column


def apply_entries(workbook, entries: list[gtc.Entry]) -> None:
    """Applied entries change the data sheets; pending ones are only highlighted with their proposal."""
    for entry in sorted(entries, key=lambda e: e.rule != "R05"):  # clear stray cells before adding columns
        ws = workbook[entry.sheet]
        if entry.applied:
            _write_value(ws, entry)
            continue
        if entry.row and entry.row >= 2 and entry.field in gtr.SHEETS[entry.sheet]["columns"]:
            cell = ws.cell(entry.row, _column_index(ws, entry.field))
            cell.fill = PENDING_FILL
            proposal = f"proposed {_display(entry.proposed_value)!r}" if entry.proposed_value is not None else "no value"
            note = f"{entry.change_id} ({entry.rule}) {entry.status}, not applied: {proposal}"
            cell.comment = Comment(note if cell.comment is None else f"{cell.comment.text}\n{note}", GENERATOR)
    for sheet, spec in gtr.SHEETS.items():
        ws = workbook[sheet]
        flag_col = _ensure_column(ws, "review_flags")
        ids_col = _ensure_column(ws, "change_ids")
        by_row: dict[int, list[gtc.Entry]] = {}
        for entry in entries:
            if entry.sheet == sheet and entry.row and entry.row >= 2:
                by_row.setdefault(entry.row, []).append(entry)
        for row, row_entries in by_row.items():
            flags = set()
            for e in row_entries:
                if e.applied and (e.rule == "R02" or e.field == "(row)"):
                    flags.add("REMOVED_DUPLICATE" if e.rule == "R02" else "REMOVED_RECLASSIFIED")
                elif e.applied and e.field == "(new row)":
                    flags.add("NEW_RECORD")
                elif not e.applied and e.status == "Pending" and e.change_type != "informational":
                    flags.add("PENDING_REVIEW")
            ws.cell(row, flag_col).value = ", ".join(sorted(flags)) or None
            ws.cell(row, ids_col).value = ", ".join(e.change_id for e in row_entries)


def add_change_log(workbook, entries: list[gtc.Entry]) -> None:
    ws = workbook.create_sheet("Change Log", 0)
    ws.append(LOG_COLUMNS + EXTRA_COLUMNS)
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP
    for e in entries:
        ai = e.ai or {}
        rule_type, reason, _ = gtr.RULES[e.rule]
        evidence = _evidence_text(ai) if ai else ""
        reviewer = ai.get("reviewer") or {}
        ws.append([
            e.change_id, e.incident_id, e.sheet, e.record_id, e.row, e.field, _display(e.original_value),
            _display(e.proposed_value), e.change_type, e.rule, gtr.RULES[e.rule][2], reason, evidence,
            ai.get("confidence", ""), "", e.status, e.reviewer, e.notes,
            "yes" if e.applied else "no", _display(e.applied_value) if e.applied else "", e.verification,
            ai.get("decision", ""), ai.get("confidence", ""), ai.get("explanation", ""),
            f"{reviewer.get('model', '')} / {ai.get('correction_prompt', '')}" if ai else "",
            "; ".join(e.problems), _display(e.accepted_value) if e.accepted_value is not None else None,
            e.carried_from,
        ])
        ws.cell(ws.max_row, 9).fill = FILLS[rule_type]
        if e.applied:
            ws.cell(ws.max_row, 19).fill = FILLS["verified_correction"]
    status = DataValidation(type="list", formula1='"' + ",".join(gtc.STATUSES) + '"', allow_blank=False,
                            showErrorMessage=True, errorTitle="Review status", error=", ".join(gtc.STATUSES))
    ws.add_data_validation(status)
    status.add(f"P2:P{max(ws.max_row, 2)}")
    for index, width in enumerate([10, 16, 15, 12, 6, 22, 45, 45, 20, 6, 18, 40, 50, 10, 6, 13, 14, 30,
                                   9, 30, 40, 14, 10, 50, 30, 30, 30, 18], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def add_summary(workbook, entries: list[gtc.Entry], meta: dict) -> None:
    ws = workbook.create_sheet("Review Summary", 0)
    last = len(entries) + 1
    rng = lambda col: f"'Change Log'!${col}$2:${col}${last}"  # noqa: E731
    ws["A1"], ws["A1"].font = f"Ground Truth Candidate {meta['version']}", Font(bold=True, size=14)
    info = [
        ("Status", "CANDIDATE, not ground truth. Only Auto-verified and Approved entries are applied to the data "
                   "sheets; Pending proposals are shown, not applied. PR 5 freezes after the final annotations are "
                   "approved."),
        ("Source workbook", f"{meta['source_name']} ({meta['source_sha256']})"),
        ("Decisions carried from", f"{meta['prior_name']} ({meta['prior_sha256']})"),
        ("AI proposals", f"{meta['proposals_name']} ({meta['proposals_sha256']})"),
        ("Correction prompt", f"{gtc.CORRECTION_PROMPT_VERSION} ({meta['prompt_sha256']})"),
        ("Guidelines", gtr.GUIDELINE_VERSION),
        ("Generator", f"{GENERATOR} at git {meta['git_sha']}"),
        ("How to review", "In Change Log set Review status to Approved (apply the proposed value, or the Accepted "
                          "value when you fill it in) or Rejected (keep the original). For a needs_video_review entry "
                          "with no proposed value, Approved confirms the original. Rebuild a new version to apply."),
    ]
    for offset, (label, value) in enumerate(info, start=3):
        ws.cell(offset, 1, label).font = Font(bold=True)
        ws.cell(offset, 2, value).alignment = WRAP
    row = len(info) + 5
    ws.cell(row, 1, "Progress").font = Font(bold=True, size=12)
    for status in ("Total", *gtc.STATUSES):
        row += 1
        ws.cell(row, 1, status)
        ws.cell(row, 2, f"=COUNTA({rng('A')})" if status == "Total" else f'=COUNTIF({rng("P")},"{status}")')
    row += 1
    ws.cell(row, 1, "Applied in this candidate")
    ws.cell(row, 2, f'=COUNTIF({rng("S")},"yes")')
    row += 2
    ws.cell(row, 1, "By rule").font = Font(bold=True, size=12)
    row += 1
    for column, title in enumerate(["Rule", "Class", "Entries", "Pending", "Applied", "AI: change / keep / uncertain"],
                                   start=1):
        ws.cell(row, column, title).font = Font(bold=True)
    for rule in sorted({e.rule for e in entries}):
        row += 1
        group = [e for e in entries if e.rule == rule]
        ai = Counter((e.ai or {}).get("decision") for e in group)
        ws.cell(row, 1, rule)
        ws.cell(row, 2, gtr.RULES[rule][0]).fill = FILLS[gtr.RULES[rule][0]]
        ws.cell(row, 3, f'=COUNTIF({rng("J")},"{rule}")')
        ws.cell(row, 4, f'=COUNTIFS({rng("J")},"{rule}",{rng("P")},"Pending")')
        ws.cell(row, 5, f'=COUNTIFS({rng("J")},"{rule}",{rng("S")},"yes")')
        ws.cell(row, 6, f"{ai['propose_change']} / {ai['keep_original']} / {ai['uncertain']}" if any(ai.values())
                else "")
    for column, width in zip("ABCDEF", (26, 90, 10, 10, 10, 28), strict=True):
        ws.column_dimensions[column].width = width


def report(entries: list[gtc.Entry], meta: dict, problems: list[str]) -> str:
    by_rule: dict[str, list[gtc.Entry]] = {}
    for e in entries:
        by_rule.setdefault(e.rule, []).append(e)
    status = Counter(e.status for e in entries)
    lines = [
        f"# Ground-truth candidate {meta['version']}",
        "",
        "> Generated by `eval/scripts/gt_candidate_build.py`. **A candidate, not ground truth.** Only `Auto-verified` "
        "(independently re-checked mechanical corrections) and `Approved` (human) entries are applied; AI video-review "
        "proposals stay `Pending` until a person decides them. Ids and counts only; the text is in the workbook.",
        "",
        "## Inputs",
        "",
        "| Item | Value |",
        "|---|---|",
        f"| Source workbook | `{meta['source_name']}` `{meta['source_sha256']}` (unchanged) |",
        f"| Decisions carried from | `{meta['prior_name']}` `{meta['prior_sha256']}` (unchanged): "
        f"{meta['prior_state']['statuses']} |",
        f"| AI proposals | `{meta['proposals_name']}` `{meta['proposals_sha256']}`: {meta['proposals']} lines |",
        f"| Correction prompt | `{gtc.CORRECTION_PROMPT_VERSION}` `{meta['prompt_sha256']}` |",
        f"| Guidelines | `{gtr.GUIDELINE_VERSION}` |",
        f"| Entries digest | `{meta['entries_sha256']}` |",
        f"| Corrected records digest | `{meta['records_sha256']}` |",
        "",
        "## Status",
        "",
        f"{len(entries)} entries: " + ", ".join(f"{status[s]} {s}" for s in gtc.STATUSES) + f". Applied: "
        f"{sum(e.applied for e in entries)}.",
        "",
        "| Rule | Class | Entries | Applied | Pending | AI change | AI keep | AI uncertain | AI confidence h/m/l | "
        "Problems |",
        "|---|---|---|---|---|---|---|---|---|---|",
    ]
    for rule in sorted(by_rule):
        group = by_rule[rule]
        ai = Counter((e.ai or {}).get("decision") for e in group)
        conf = Counter((e.ai or {}).get("confidence") for e in group)
        lines.append(
            f"| {rule} | {gtr.RULES[rule][0]} | {len(group)} | {sum(e.applied for e in group)} | "
            f"{sum(e.status == 'Pending' for e in group)} | {ai['propose_change']} | {ai['keep_original']} | "
            f"{ai['uncertain']} | {conf['high']}/{conf['medium']}/{conf['low']} | "
            f"{sum(bool(e.problems) for e in group)} |")
    lines += ["", "## Checks", ""]
    lines += [f"- {p}" for p in problems] or ["- No consistency or carry-forward problems."]
    lines += [
        f"- Mechanical corrections failing the independent re-check (left Pending): "
        f"{', '.join(e.change_id for e in entries if e.change_type == 'verified_correction' and e.status == 'Pending') or 'none'}.",
        f"- Findings without an AI proposal: "
        f"{', '.join(e.change_id for e in entries if e.change_type in ('needs_video_review', 'rule_based_proposal') and not e.ai) or 'none'}.",
        f"- AI proposals with validation problems (shown, never usable): "
        f"{', '.join(e.change_id for e in entries if e.problems) or 'none'}.",
        "",
        "## Next",
        "",
        "A person reviews the `Pending` entries in the candidate workbook (every AI proposal and both rule-based "
        "proposals), then a new candidate version is built from that workbook. PR 5 freezes `ground-truth-v1.0` only "
        "after the final annotations and exclusions are explicitly approved.",
        "",
    ]
    return "\n".join(lines)


def build(source: Path, manifest_path: Path, prior: Path, proposals_path: Path, out_dir: Path, version: str, *,
          overwrite: bool = False, frames_dir: Path = FRAMES_DIR) -> dict:
    output = out_dir / f"Ground Truth Candidate {version}.xlsx"
    if output.resolve() in (source.resolve(), prior.resolve()):
        raise review_backup.OverwriteRefusedError("the candidate must be a new file, not an input")
    gtc.guard_inputs([source, manifest_path, prior, proposals_path, frames_dir])
    previous_state = review_backup.guard_overwrite(output, overwrite=overwrite, discard_review_decisions=False)
    hashes = {path: sha256_file(path) for path in (source, prior)}
    manifest = json.loads(manifest_path.read_text())
    values = openpyxl.load_workbook(source, data_only=True)
    changes, _ = gtr.audit(values, manifest)
    prior_log = gtc.read_change_log(openpyxl.load_workbook(prior, data_only=True))
    problems = gtc.check_carry_forward(changes, prior_log)
    if problems:
        raise ValueError(f"{prior.name} does not match the audit: {problems[:3]}")
    prior_by_id = {str(row["Change ID"]): row for row in prior_log}

    verification = {c.change_id: gtc.verify_mechanical(c, values, manifest) for c in changes
                    if c.change_type == "verified_correction"}
    proposals = gtc.load_proposals(proposals_path)
    by_id = {c.change_id: c for c in changes}
    proposal_problems = {}
    for change_id, proposal in proposals.items():
        change = by_id.get(change_id)
        if change is None:
            proposal_problems[change_id] = ["no such change"]
            continue
        frames_file = frames_dir / change.incident_id / "frames.json"
        frames = json.loads(frames_file.read_text()) if frames_file.exists() else None
        proposal_problems[change_id] = gtc.validate_proposal(proposal, change, frames)
    entries = gtc.resolve(changes, prior=prior_by_id, prior_name=prior.name, verification=verification,
                          proposals=proposals, proposal_problems=proposal_problems)
    problems = gtc.consistency_problems(entries)
    problems += [f"{cid}: {p}" for cid, ps in proposal_problems.items() if cid not in by_id for p in ps]
    records = gtc.apply_to_records(gtr.read_data_rows(values), entries)

    working = openpyxl.load_workbook(source)  # formulas kept; saved under a new name only
    apply_entries(working, entries)
    try:
        import subprocess
        git_sha = subprocess.run(["git", "rev-parse", "--short", "HEAD"], capture_output=True, text=True, cwd=EVAL_DIR,
                                 check=False).stdout.strip() or "unknown"
    except OSError:
        git_sha = "unknown"
    meta = {
        "version": version,
        "source_name": source.name,
        "source_sha256": hashes[source],
        "prior_name": prior.name,
        "prior_sha256": hashes[prior],
        "prior_state": review_backup.review_state(prior),
        "proposals_name": proposals_path.name,
        "proposals_sha256": sha256_file(proposals_path) if proposals_path.exists() else None,
        "proposals": len(proposals),
        "prompt_sha256": sha256_file(PROMPT_PATH),
        "entries_sha256": gtc.entries_digest(entries),
        "records_sha256": gtc.records_digest(records),
        "git_sha": git_sha,
    }
    add_change_log(working, entries)
    add_summary(working, entries, meta)
    working.active = 0
    out_dir.mkdir(parents=True, exist_ok=True)
    backup_record = review_backup.backup(output, reason=f"before rebuild of {version}") if previous_state else None
    working.save(output)
    for path, digest in hashes.items():
        if sha256_file(path) != digest:
            raise RuntimeError(f"{path.name} changed during the build")
    result = {
        **meta,
        "generator": GENERATOR,
        "guidelines": gtr.GUIDELINE_VERSION,
        "correction_prompt": gtc.CORRECTION_PROMPT_VERSION,
        "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "output_name": output.name,
        "output_sha256": sha256_file(output),
        "entries": len(entries),
        "statuses": dict(sorted(Counter(e.status for e in entries).items())),
        "applied": sum(e.applied for e in entries),
        "applied_by_rule": dict(sorted(Counter(e.rule for e in entries if e.applied).items())),
        "ai_decisions": dict(sorted(Counter((e.ai or {}).get("decision", "none") for e in entries
                                            if e.change_type in ("needs_video_review", "rule_based_proposal")).items())),
        "ai_proposal_problems": sum(bool(e.problems) for e in entries),
        "problems": problems,
        "corrected_records": {sheet: len(rows) for sheet, rows in records.items()},
        "replaced_backup": backup_record["backup"] if backup_record else None,
    }
    (out_dir / f"candidate-{version}-manifest.json").write_text(json.dumps(result, indent=2, default=str) + "\n")
    (out_dir / f"candidate-{version}.md").write_text(report(entries, meta, problems))
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a ground-truth candidate workbook (PR 4).")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--prior", type=Path, default=DEFAULT_PRIOR, help="Review workbook whose decisions carry forward")
    parser.add_argument("--proposals", type=Path, help="AI proposals JSONL (default ai-proposals-<version>.jsonl)")
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR)
    parser.add_argument("--version", default="v0.2")
    parser.add_argument("--overwrite", action="store_true",
                        help="Replace an existing candidate of this version that holds no human decisions (backed up)")
    args = parser.parse_args()
    proposals = args.proposals or args.out_dir / f"ai-proposals-{args.version}.jsonl"
    try:
        result = build(args.source, args.manifest, args.prior, proposals, args.out_dir, args.version,
                       overwrite=args.overwrite)
    except review_backup.OverwriteRefusedError as exc:
        print(f"REFUSED: {exc}", file=sys.stderr)
        return 2
    print(json.dumps({k: result[k] for k in ("output_name", "output_sha256", "entries", "statuses", "applied",
                                             "ai_decisions", "ai_proposal_problems", "problems")}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
