# SPDX-FileCopyrightText: Copyright (c) 2025-2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Build the ground-truth review workbook, the audit report and the review manifest (PR 3).

The source workbook is opened twice, read-only in effect: once for values (the audit)
and once with its formulas (the base of the working copy), and it is never saved. Its
SHA-256 is checked before and after. The working copy keeps every sheet, column, record
id and relationship of the source. It writes proposed values in place (original values
go to the ``Change Log`` sheet), appends new records instead of renumbering, and marks
removals with a review flag instead of deleting rows. It also adds the ``Change Log`` and
``Review Summary`` sheets.

An existing workbook of the same version is never replaced silently: the build refuses unless
``--overwrite`` is passed, and refuses even then when the workbook holds review decisions
(a status other than Pending, a reviewer or notes) unless ``--discard-review-decisions`` is
passed too. A permitted overwrite first takes a timestamped backup (``review_backup.py``).

Usage (from ``eval/``)::

    uv run python scripts/gt_review_build.py            # defaults: the profile-root workbook, v0.1
    uv run python scripts/gt_review_build.py --source <workbook> --version v0.2

Outputs in ``ground_truth/review/``:
- ``Ground Truth Review <version>.xlsx`` - the editable working workbook (gitignored; see README.md there);
- ``audit-<version>.md`` - the audit report (ids and counts only, no annotation text);
- ``review-<version>-manifest.json`` - source and output hashes, the change-log digest and counts.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import ground_truth_review as gtr  # noqa: E402
import openpyxl  # noqa: E402
import review_backup  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.formula.translate import Translator  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

EVAL_DIR = Path(__file__).resolve().parents[1]
PROFILE_DIR = EVAL_DIR.parent
DEFAULT_SOURCE = PROFILE_DIR / "Complete Ground Truth Labelling.xlsx"
DEFAULT_MANIFEST = EVAL_DIR / "eval_reproducibility" / "dataset-manifest-v1.json"
DEFAULT_OUT_DIR = EVAL_DIR / "ground_truth" / "review"
GENERATOR = "gt_review_build.py v1"

FILLS = {
    "verified_correction": PatternFill("solid", fgColor="C6EFCE"),
    "rule_based_proposal": PatternFill("solid", fgColor="FFEB9C"),
    "needs_video_review": PatternFill("solid", fgColor="FCE4D6"),
    "informational": PatternFill("solid", fgColor="DDEBF7"),
}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
LOG_COLUMNS = ["Change ID", "Incident ID", "Sheet", "Record ID", "Row", "Field", "Original value", "Proposed value",
               "Change type", "Rule", "Guideline ref", "Reason", "Supporting video evidence", "Confidence",
               "Review flag", "Review status", "Reviewer", "Review notes"]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _display(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        return "new row: " + "; ".join(f"{k}={'' if v is None else v}" for k, v in value.items())
    return str(value)


def _column_index(ws, header: str) -> int:
    for cell in ws[1]:
        if str(cell.value or "").strip() == header:
            return cell.column
    raise KeyError(f"{ws.title}: no column {header!r}")


def _last_data_row(ws, documented: int) -> int:
    last = 1
    for row in ws.iter_rows(min_row=2, max_col=documented):
        if any(cell.value not in (None, "") and not str(cell.value).startswith("=") for cell in row):
            last = row[0].row
    return last


def apply_changes(workbook, changes: list[gtr.Change]) -> dict[tuple[str, int], list[gtr.Change]]:
    """Write proposals into the working copy; returns the changes per (sheet, row) after appending new rows."""
    by_row: dict[tuple[str, int], list[gtr.Change]] = defaultdict(list)
    # 1. Clear cells outside the schema first, so the review columns never overwrite data.
    for change in changes:
        if change.rule == "R05":
            column = int(change.field.split()[1])
            workbook[change.sheet].cell(change.row, column).value = None
    # 2. Append new records (never renumber), then edit values in place.
    for change in changes:
        ws = workbook[change.sheet]
        if change.field == "(new row)":
            previous = _last_data_row(ws, len(gtr.SHEETS[change.sheet]["columns"]))
            row = previous + 1
            for header, value in change.proposed_value.items():
                column = _column_index(ws, header)
                above = ws.cell(previous, column).value
                if isinstance(above, str) and above.startswith("="):
                    # The column is derived (e.g. Incident_ID from Filename): carry the formula, not a literal.
                    origin = f"{get_column_letter(column)}{previous}"
                    value = Translator(above, origin=origin).translate_formula(f"{get_column_letter(column)}{row}")
                ws.cell(row, column).value = value
            change.row = row
            for column in range(1, len(gtr.SHEETS[change.sheet]["columns"]) + 1):
                ws.cell(row, column).fill = FILLS[change.change_type]
        elif change.change_type in ("verified_correction", "rule_based_proposal") and change.field not in ("(row)",) \
                and change.rule != "R05":
            column = _column_index(ws, change.field)
            cell = ws.cell(change.row, column)
            if change.rule == "R28":
                cell.value = f"=E{change.row}-D{change.row}"
            else:
                cell.value = change.proposed_value
            cell.fill = FILLS[change.change_type]
            cell.comment = Comment(f"{change.change_id} ({change.rule}): was {_display(change.original_value)!r}",
                                   "gt_review_build")
        elif change.change_type == "needs_video_review" and change.field in gtr.SHEETS[change.sheet]["columns"]:
            cell = ws.cell(change.row, _column_index(ws, change.field))
            cell.fill = FILLS["needs_video_review"]
            note = f"{change.change_id} ({change.rule}): needs video review"
            cell.comment = Comment(note if cell.comment is None else f"{cell.comment.text}\n{note}", "gt_review_build")
        by_row[(change.sheet, change.row)].append(change)
    # 3. Review columns right of the documented ones.
    for sheet, spec in gtr.SHEETS.items():
        ws = workbook[sheet]
        documented = len(spec["columns"])
        flag_col, ids_col = documented + 1, documented + 2
        for column, title in ((flag_col, "review_flags"), (ids_col, "change_ids")):
            header = ws.cell(1, column)
            header.value, header.font = title, Font(bold=True)
        for (row_sheet, row), row_changes in by_row.items():
            if row_sheet != sheet or row < 2:
                continue
            flags = sorted({c.review_flag for c in row_changes if c.review_flag}
                           | ({"NEEDS_VIDEO_REVIEW"} if any(c.change_type == "needs_video_review" for c in row_changes)
                              else set()))
            ws.cell(row, flag_col).value = ", ".join(flags) or None
            ws.cell(row, ids_col).value = ", ".join(c.change_id for c in row_changes)
            if any(c.review_flag.startswith("PROPOSED_REMOVAL") for c in row_changes):
                for column in range(1, documented + 1):
                    ws.cell(row, column).font = Font(strike=True, color="7F7F7F")
    return by_row


def add_change_log(workbook, changes: list[gtr.Change]) -> None:
    ws = workbook.create_sheet("Change Log", 0)
    ws.append(LOG_COLUMNS)
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP
    for change in changes:
        ws.append([change.change_id, change.incident_id, change.sheet, change.record_id, change.row, change.field,
                   _display(change.original_value), _display(change.proposed_value), change.change_type, change.rule,
                   change.guideline_ref, change.reason, change.evidence, change.confidence, change.review_flag,
                   change.review_status, None, None])
        ws.cell(ws.max_row, 9).fill = FILLS[change.change_type]
    status = DataValidation(type="list", formula1='"' + ",".join(gtr.REVIEW_STATUSES) + '"', allow_blank=False,
                            showErrorMessage=True, errorTitle="Review status", error="Pending, Approved or Rejected")
    ws.add_data_validation(status)
    status.add(f"P2:P{max(ws.max_row, 2)}")
    widths = [10, 16, 15, 12, 6, 22, 45, 45, 20, 6, 18, 50, 50, 20, 26, 12, 14, 30]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LOG_COLUMNS))}{ws.max_row}"


def add_review_summary(workbook, changes: list[gtr.Change], meta: dict) -> None:
    ws = workbook.create_sheet("Review Summary", 0)
    log = "'Change Log'"
    last = len(changes) + 1
    rng = lambda col: f"{log}!${col}$2:${col}${last}"  # noqa: E731
    ws["A1"], ws["A1"].font = "Ground Truth Review " + meta["version"], Font(bold=True, size=14)
    info = [
        ("Status", "Proposed corrections for review. NOT final ground truth: PR 4 reviews, PR 5 freezes."),
        ("Source workbook", meta["source_name"]),
        ("Source SHA-256", meta["source_sha256"]),
        ("Guidelines", gtr.GUIDELINE_VERSION),
        ("Dataset manifest", f"dataset-manifest-v1 (clips_sha256 {meta['clips_sha256'][:16]}…)"),
        ("Generator", f"{GENERATOR} at git {meta['git_sha']}"),
        ("Change-log digest", meta["change_log_sha256"]),
        ("How to review", "Set 'Review status' in Change Log to Approved or Rejected; add Reviewer and notes. "
                          "This sheet updates from those cells."),
    ]
    for offset, (label, value) in enumerate(info, start=3):
        ws.cell(offset, 1, label).font = Font(bold=True)
        ws.cell(offset, 2, value).alignment = WRAP

    row = len(info) + 5
    ws.cell(row, 1, "Progress").font = Font(bold=True, size=12)
    row += 1
    for label, formula in (
        ("Total entries", f"=COUNTA({rng('A')})"),
        ("Pending", f'=COUNTIF({rng("P")},"Pending")'),
        ("Approved", f'=COUNTIF({rng("P")},"Approved")'),
        ("Rejected", f'=COUNTIF({rng("P")},"Rejected")'),
        ("Reviewed share", f"=IF(B{row}=0,0,(B{row + 2}+B{row + 3})/B{row})"),
    ):
        ws.cell(row, 1, label)
        ws.cell(row, 2, formula)
        row += 1
    ws.cell(row - 1, 2).number_format = "0%"

    row += 1
    ws.cell(row, 1, "Proposed changes and issues by class").font = Font(bold=True, size=12)
    row += 1
    for column, title in enumerate(["Change type", "Entries", "Pending", "Approved", "Rejected", "Meaning"], start=1):
        ws.cell(row, column, title).font = Font(bold=True)
    meanings = {
        "verified_correction": "Provable without video (manifest or workbook); value proposed",
        "rule_based_proposal": "Follows from an approved guideline rule; value proposed",
        "needs_video_review": "Unresolved: only the video can answer; no value proposed",
        "informational": "Recorded only; no change proposed",
    }
    for change_type in gtr.CHANGE_TYPES:
        row += 1
        ws.cell(row, 1, change_type).fill = FILLS[change_type]
        ws.cell(row, 2, f'=COUNTIF({rng("I")},"{change_type}")')
        for column, status in ((3, "Pending"), (4, "Approved"), (5, "Rejected")):
            ws.cell(row, column, f'=COUNTIFS({rng("I")},"{change_type}",{rng("P")},"{status}")')
        ws.cell(row, 6, meanings[change_type])

    row += 2
    ws.cell(row, 1, "By rule").font = Font(bold=True, size=12)
    row += 1
    for column, title in enumerate(["Rule", "Class", "Entries", "Pending", "Description", "Guideline"], start=1):
        ws.cell(row, column, title).font = Font(bold=True)
    for rule in sorted({c.rule for c in changes}):
        change_type, description, reference = gtr.RULES[rule]
        row += 1
        ws.cell(row, 1, rule)
        ws.cell(row, 2, change_type).fill = FILLS[change_type]
        ws.cell(row, 3, f'=COUNTIF({rng("J")},"{rule}")')
        ws.cell(row, 4, f'=COUNTIFS({rng("J")},"{rule}",{rng("P")},"Pending")')
        ws.cell(row, 5, description)
        ws.cell(row, 6, reference)

    row += 2
    ws.cell(row, 1, "Affected videos").font = Font(bold=True, size=12)
    row += 1
    for column, title in enumerate(["Incident ID", "Entries", "Pending", "Needs video review", "Review flags"], start=1):
        ws.cell(row, column, title).font = Font(bold=True)
    per_incident: dict[str, list[gtr.Change]] = defaultdict(list)
    for change in changes:
        per_incident[change.incident_id].append(change)
    for incident_id in sorted(per_incident):
        row += 1
        ws.cell(row, 1, incident_id)
        ws.cell(row, 2, f'=COUNTIF({rng("B")},"{incident_id}")')
        ws.cell(row, 3, f'=COUNTIFS({rng("B")},"{incident_id}",{rng("P")},"Pending")')
        ws.cell(row, 4, f'=COUNTIFS({rng("B")},"{incident_id}",{rng("I")},"needs_video_review",{rng("P")},"Pending")')
        ws.cell(row, 5, ", ".join(sorted({c.review_flag for c in per_incident[incident_id] if c.review_flag})) or None)
    for column, width in zip("ABCDEF", (24, 70, 12, 18, 60, 22), strict=True):
        ws.column_dimensions[column].width = width


def audit_report(changes: list[gtr.Change], facts: dict, meta: dict) -> str:
    by_rule: dict[str, list[gtr.Change]] = defaultdict(list)
    for change in changes:
        by_rule[change.rule].append(change)
    types = Counter(c.change_type for c in changes)
    lines = [
        f"# Ground-truth audit {meta['version']}",
        "",
        f"> Generated by `eval/scripts/gt_review_build.py` against `{gtr.GUIDELINE_VERSION}`. Proposed corrections only: "
        "nothing here is final ground truth (PR 4 reviews, PR 5 freezes). The report names ids and counts, never "
        "annotation text; the text is in the working workbook.",
        "",
        "## Sources",
        "",
        "| Item | Value |",
        "|---|---|",
        f"| Source workbook | `{meta['source_name']}` (unchanged; SHA-256 checked before and after the build) |",
        f"| Source SHA-256 | `{meta['source_sha256']}` |",
        f"| Dataset manifest | `eval_reproducibility/dataset-manifest-v1.json`, clips_sha256 `{meta['clips_sha256']}` |",
        f"| Guidelines | `{gtr.GUIDELINE_VERSION}` |",
        f"| Change-log digest | `{meta['change_log_sha256']}` (the same inputs always give this digest) |",
        f"| Working workbook | `{meta['output_name']}` (see [README](README.md) for storage) |",
        "",
        "## Inventory",
        "",
        f"- Incidents: {facts['incidents']} distinct ids in {facts['incident_rows']} incident rows; "
        f"entities {facts['rows']['gt_entities']}, instruments {facts['rows']['gt_instruments']}, "
        f"assets {facts['rows']['gt_assets']}.",
        f"- Video matching against the dataset manifest ({facts['dataset_clips']} clips): incidents without a clip: "
        f"{', '.join(facts['incidents_without_clip']) or 'none'}; clips without an incident: "
        f"{', '.join(facts['clips_without_incident']) or 'none'}.",
        f"- Type labels used: {', '.join(f'{k} ({v})' for k, v in facts['type_labels'].items())}. All are among the five "
        "categories; `Animal` is the workbook's label for `animal attack` (eval maps it), so no relabelling is proposed.",
        f"- The workbook's Type dropdown lists {facts['type_dropdown']}. `Fighting` (retired) and `Shooting` (not one of "
        "the five) are offered but used by no row; the dropdown is left as it is.",
        "",
        "## Findings",
        "",
        f"{len(changes)} entries: " + ", ".join(f"{types[t]} {t}" for t in gtr.CHANGE_TYPES) + ".",
        "",
        "| Rule | Class | Entries | What | Guideline | Records |",
        "|---|---|---|---|---|---|",
    ]
    for rule in sorted(by_rule):
        change_type, description, reference = gtr.RULES[rule]
        records = sorted({f"{c.incident_id}" + (f" {c.record_id}" if c.record_id != c.incident_id else "")
                          for c in by_rule[rule]})
        shown = ", ".join(records[:12]) + (f", … (+{len(records) - 12})" if len(records) > 12 else "")
        lines.append(f"| {rule} | {change_type} | {len(by_rule[rule])} | {description} | {reference} | {shown} |")
    lines += [
        "",
        "## How the classes are used",
        "",
        "- **verified_correction:** provable without the video, from the workbook or the dataset manifest. The value is "
        "written into the working workbook (green).",
        "- **rule_based_proposal:** follows from an approved rule applied to what the annotator recorded. The value is "
        "written (yellow), and a video check is recommended.",
        "- **needs_video_review:** only the video can answer, so **no value is proposed**; the cell is marked (orange). "
        "Description wording only selects records for review here; it never supplies a value.",
        "- **informational:** recorded, no change.",
        "- Rows proposed for removal stay in the sheet, struck through, with a review flag; new records are appended with "
        "the next free id. Nothing is renumbered.",
        "- Possible out-of-taxonomy clips carry the `OUT_OF_TAXONOMY_CANDIDATE` flag; their type stays one of the five.",
        "",
        "## Not changed",
        "",
        "- The 9 incidents without entities (A2: unseen drivers are not entities).",
        "- Image cells holding links: image fields are intentionally empty or informational (guidelines §9).",
        "- Severity and threat values stored as whole-number floats (`3.0`): equal to the integers, not an error.",
        "",
    ]
    return "\n".join(lines)


def build(source: Path, manifest_path: Path, out_dir: Path, version: str, *, overwrite: bool = False,
          discard_review_decisions: bool = False) -> dict:
    """Build one version. An existing workbook of that version is never replaced silently.

    The guard runs before any work (``review_backup.guard_overwrite``). A refused build
    writes nothing; a permitted overwrite backs the old workbook up first.
    """
    output = out_dir / f"Ground Truth Review {version}.xlsx"
    previous_state = review_backup.guard_overwrite(output, overwrite=overwrite,
                                                   discard_review_decisions=discard_review_decisions)
    source_hash = sha256_file(source)
    manifest = json.loads(manifest_path.read_text())
    values = openpyxl.load_workbook(source, data_only=True)
    changes, facts = gtr.audit(values, manifest)
    working = openpyxl.load_workbook(source)  # formulas kept; saved under a new name only
    apply_changes(working, changes)
    try:
        git_sha = subprocess.run(["git", "rev-parse", "--short", "HEAD"], capture_output=True, text=True,
                                 cwd=EVAL_DIR, check=False).stdout.strip() or "unknown"
    except OSError:
        git_sha = "unknown"
    out_dir.mkdir(parents=True, exist_ok=True)
    meta = {
        "version": version,
        "source_name": source.name,
        "source_sha256": source_hash,
        "clips_sha256": manifest["clips_sha256"],
        "change_log_sha256": gtr.change_log_digest(changes),
        "git_sha": git_sha,
        "output_name": output.name,
    }
    add_change_log(working, changes)
    add_review_summary(working, changes, meta)
    working.active = 0
    backup_record = None
    if previous_state is not None:
        backup_record = review_backup.backup(output, reason=f"before rebuild of {version}")
    working.save(output)
    if sha256_file(source) != source_hash:
        raise RuntimeError("the source workbook changed during the build")
    review_manifest = {
        **meta,
        "generator": GENERATOR,
        "guidelines": gtr.GUIDELINE_VERSION,
        "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "output_sha256": sha256_file(output),
        "entries": len(changes),
        "entries_by_class": dict(Counter(c.change_type for c in changes)),
        "entries_by_rule": dict(sorted(Counter(c.rule for c in changes).items())),
        "affected_incidents": len({c.incident_id for c in changes}),
        "replaced_backup": backup_record["backup"] if backup_record else None,
        "facts": facts,
    }
    (out_dir / f"review-{version}-manifest.json").write_text(json.dumps(review_manifest, indent=2, default=str) + "\n")
    (out_dir / f"audit-{version}.md").write_text(audit_report(changes, facts, meta))
    return review_manifest


def main() -> int:
    parser = argparse.ArgumentParser(description="Build the ground-truth review workbook and audit report.")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR)
    parser.add_argument("--version", default="v0.1")
    parser.add_argument("--overwrite", action="store_true",
                        help="Replace an existing workbook of this version (it is backed up first).")
    parser.add_argument("--discard-review-decisions", action="store_true",
                        help="With --overwrite: allow replacing a workbook that holds review decisions (backed up first).")
    args = parser.parse_args()
    if args.discard_review_decisions and not args.overwrite:
        parser.error("--discard-review-decisions requires --overwrite")
    try:
        result = build(args.source, args.manifest, args.out_dir, args.version, overwrite=args.overwrite,
                       discard_review_decisions=args.discard_review_decisions)
    except review_backup.OverwriteRefusedError as exc:
        print(f"REFUSED: {exc}", file=sys.stderr)
        return 2
    print(json.dumps({k: result[k] for k in ("output_name", "source_sha256", "output_sha256", "change_log_sha256",
                                             "entries", "entries_by_class", "affected_incidents")}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
