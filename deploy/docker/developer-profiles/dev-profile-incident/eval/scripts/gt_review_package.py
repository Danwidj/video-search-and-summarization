# SPDX-FileCopyrightText: Copyright (c) 2025-2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Build the human review package for the Pending entries of a candidate workbook (PR 4 follow-up).

Inputs, none of which is ever written: the source workbook (for the traceability checks), the
candidate workbook (``Ground Truth Candidate <v>.xlsx``), its AI proposals, the contact sheets,
an optional earlier package whose decisions are carried forward, and an optional local overrides
file (``review-overrides-<v>.json``, gitignored: it quotes annotation text) with suggested pronoun
fixes and case-specific recommendations.

Outputs, all new files (an existing one is never overwritten):
- ``ground_truth/review/Ground Truth Review Package <version>.xlsx``: where decisions are recorded.
  Its ``Change Log`` holds every entry (Pending rows grouped by incident and correction type,
  then the Auto-verified rows, greyed) with the columns ``gt_candidate_build.py --prior`` reads,
  so the next candidate is built from it directly;
- ``eval_data/gt_review_package/review-package-<version>.html``: a read-only companion with the
  cited evidence tiles cut from the contact sheets;
- ``eval_data/gt_review_package/verification-<version>.json``: the traceability checks.

Pronoun policy P1 (decisions.md 2026-10-27): a gendered pronoun may follow a visually supported,
qualified gender identification ("appears to be a man"); where gender cannot be determined and the
wording is neutral, later pronouns must be neutral too. Every rewrite that keeps a gendered pronoun
is flagged; the overrides name the definite violations and their suggested correction.

No entry is ever approved here: every Pending row stays Pending.

Usage (from ``eval/``)::

    uv run python scripts/gt_review_package.py --candidate "ground_truth/review/Ground Truth Candidate v0.2.xlsx" \\
        --version v0.2.1 [--prior-package "ground_truth/review/Ground Truth Review Package v0.2.xlsx"]
"""

from __future__ import annotations

import argparse
import base64
import html
import io
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import ground_truth_correction as gtc  # noqa: E402
import ground_truth_review as gtr  # noqa: E402
import openpyxl  # noqa: E402
import review_backup  # noqa: E402
from gt_review_build import DEFAULT_OUT_DIR, DEFAULT_SOURCE, EVAL_DIR, sha256_file  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

GENERATOR = "gt_review_package.py v1"
FRAMES_DIR = EVAL_DIR / "eval_data" / "gt_frames"
HTML_DIR = EVAL_DIR / "eval_data" / "gt_review_package"
TYPES = {"R30": "Informational (ids)", "R01": "Mechanical", "R02": "Mechanical", "R03": "Mechanical",
         "R05": "Mechanical", "R28": "Mechanical", "R25": "Age/gender wording", "R23": "Animal severity", "R20": "Explosive threat level",
         "R21": "Object name", "R22": "Out-of-taxonomy check", "R24": "Responder equipment",
         "R34": "Unobservable conclusion", "R10": "Vehicle reclassified as asset"}
HIGH_IMPACT_RULES = ("R23", "R20", "R22", "R10")
PRONOUN = re.compile(r"\b(he|him|his|himself|she|her|hers|herself)\b", re.IGNORECASE)
# The Change Log columns. Review status, Reviewer and Review notes stay columns 16-18 (review_backup).
COLUMNS = ["Change ID", "Incident ID", "Correction type", "Sheet", "Record ID", "Row", "Field", "Rule", "Before",
           "After (AI proposal)", "AI decision", "AI confidence", "Evidence (tile time: observation)",
           "Recommendation", "Flags", "Review status", "Reviewer", "Review notes", "Accepted value",
           "Suggested accepted value", "Why", "Frames"]
HUMAN_COLUMNS = ("Review status", "Reviewer", "Review notes", "Accepted value")


def recommend(rule: str, decision: str, confidence: str, violation: bool) -> tuple[str, str]:
    """The general recommendation; case-specific ones come from the overrides file."""
    if decision == "keep_original":
        return ("Approve (confirms original)",
                "The AI found the annotation supported. Approving records the check and changes nothing.")
    if decision == "propose_change" and violation:
        return ("Approve with Accepted value (P1 pronoun fix)",
                "The rewrite makes a person gender-neutral but a later pronoun still genders them (P1). Use the "
                "suggested accepted value.")
    if decision == "propose_change" and rule == "R25":
        if confidence == "low":
            return ("Approve (conservative under A7)",
                    "Only the unsupported age or gender words change. Low confidence means the cue is weak, and A7 "
                    "asks for clear cues, so the hedged or neutral wording is the conservative choice.")
        return "Approve", "Only the unsupported age or gender words change."
    if decision == "propose_change":
        return "Approve", "Follows the cited guideline rule. Check the evidence tiles."
    if decision == "none":
        return ("No action (informational)" if rule == "R30" else "Check the mechanical fix",
                "Not an AI video-review entry.")
    return "Watch, then decide", "The AI could not settle it."


def pronoun_flag(change_id: str, decision: str, after, overrides: dict) -> str | None:
    if decision != "propose_change" or not isinstance(after, str):
        return None
    if change_id in overrides.get("pronoun_fixes", {}):
        return "PRONOUN: P1 violation (correct)"
    words = sorted({w.lower() for w in PRONOUN.findall(after)})
    if words:
        return f"PRONOUN: permitted under P1 if the qualified identification is visually supported ({', '.join(words)})"
    return None


def suggested_value(change_id: str, after, original, overrides: dict):
    if change_id in overrides.get("pronoun_fixes", {}):
        value = str(after)
        for old, new in overrides["pronoun_fixes"][change_id]:
            if old not in value:
                raise ValueError(f"{change_id}: pronoun fix {old!r} is not in the proposed text")
            value = value.replace(old, new)
        return value
    if change_id in overrides.get("original_rewrites", {}):
        value = str(original)
        for pattern, replacement in overrides["original_rewrites"][change_id]:
            value = re.sub(pattern, replacement, value)
        return value
    return None


def build_rows(log: list[dict], proposals: dict, overrides: dict, prior: dict[str, dict]) -> list[dict]:
    rows = []
    for entry in log:
        if entry["Review status"] in ("Auto-verified", "Informational"):
            continue
        change_id = entry["Change ID"]
        proposal = proposals.get(change_id)
        if proposal is None:
            # Informational entries, and mechanical fixes that failed their re-check, have no AI proposal.
            kind = entry["Change type"]
            proposal = {"decision": "none", "confidence": "", "proposed_value": entry["Proposed value"] or None,
                        "evidence": [], "explanation": f"No AI proposal ({kind}; {entry['Independent verification']})."}
        decision, confidence = proposal["decision"], proposal["confidence"]
        after = proposal["proposed_value"]
        pronoun = pronoun_flag(change_id, decision, after, overrides)
        flags = [f for f in ("UNCERTAIN" if decision == "uncertain" else None, pronoun,
                             "HIGH IMPACT" if entry["Rule"] in HIGH_IMPACT_RULES else None) if f]
        recommendation, why = overrides.get("recommendations", {}).get(change_id) or recommend(
            entry["Rule"], decision, confidence, bool(pronoun and "violation" in pronoun))
        carried = prior.get(change_id, {})
        rows.append({
            "Change ID": change_id, "Incident ID": entry["Incident ID"], "Correction type": TYPES.get(entry["Rule"], gtr.RULES[entry["Rule"]][1]),
            "Sheet": entry["Sheet"], "Record ID": entry["Record ID"], "Row": entry["Row"], "Field": entry["Field"],
            "Rule": entry["Rule"], "Before": entry["Original value"],
            "After (AI proposal)": json.dumps(after) if isinstance(after, dict) else after
            if decision in ("propose_change", "none") else ("(no change: original confirmed)"
                                                             if decision == "keep_original"
                                                             else "(no proposal: uncertain)"),
            "AI decision": decision, "AI confidence": confidence,
            "Evidence (tile time: observation)": "; ".join(
                f"{e['sheet']} @{e['t']}s: {e.get('observation', '')}" for e in proposal["evidence"]) or "(none)",
            "Recommendation": recommendation, "Flags": "; ".join(flags),
            # A decision is only ever carried from an earlier package; nothing is decided here.
            "Review status": carried.get("Review status") or entry["Review status"],
            "Reviewer": carried.get("Reviewer"), "Review notes": carried.get("Review notes"),
            "Accepted value": carried.get("Accepted value"),
            "Suggested accepted value": suggested_value(change_id, after, entry["Original value"], overrides),
            "Why": f"{why} AI: {proposal['explanation']}", "Frames": f"eval_data/gt_frames/{entry['Incident ID']}/",
            "_evidence": proposal["evidence"],
        })
    order = {t: i for i, t in enumerate(TYPES.values())}
    rows.sort(key=lambda r: (r["Incident ID"], order.get(r["Correction type"], len(order)), r["Change ID"]))
    return rows


def verification(source_path: Path, candidate, log: list[dict], proposals: dict, freeze_manifest: Path | None) -> dict:
    """The 86-style traceability checks: mechanical fixes traced and applied, nothing Pending applied."""
    pending = [r for r in log if r["Review status"] == "Pending"]
    auto = [r for r in log if r["Review status"] == "Auto-verified"]
    source = openpyxl.load_workbook(source_path, data_only=True)
    traced = untraced = still_original = changed = 0
    for sheet in gtr.SHEETS:
        ws = candidate[sheet]
        header = [str(c.value or "").strip() for c in ws[1]]
        ids_col = header.index("change_ids") + 1
        for r in auto:
            if r["Sheet"] == sheet and r["Row"]:
                if r["Change ID"] in str(ws.cell(int(r["Row"]), ids_col).value or ""):
                    traced += 1
                else:
                    untraced += 1
    for r in pending:
        if r["Field"] in gtr.SHEETS[r["Sheet"]]["columns"] and r["Row"]:
            column = gtr.SHEETS[r["Sheet"]]["columns"].index(r["Field"]) + 1
            same = candidate[r["Sheet"]].cell(int(r["Row"]), column).value == \
                source[r["Sheet"]].cell(int(r["Row"]), column).value
            still_original += same
            changed += not same
    checks = {
        "entries": len(log), "statuses": dict(Counter(r["Review status"] for r in log)),
        "auto_verified_applied_and_passed": all(r["Applied in candidate"] == "yes" and
                                                str(r["Independent verification"]).startswith("passed") for r in auto),
        "auto_verified_by_rule": dict(Counter(r["Rule"] for r in auto)),
        "auto_verified_traced_in_data_sheets": traced, "auto_verified_not_traced": untraced,
        "pending_applied": sum(r["Applied in candidate"] == "yes" for r in pending),
        "filename_changes_applied": sum(r["Applied in candidate"] == "yes" for r in log if r["Field"] == "Filename"),
        "pending_with_ai_proposal": sum(r["Change ID"] in proposals for r in pending),
        "pending_cells_still_original": still_original, "pending_cells_changed": changed,
        "human_decisions_in_candidate": sum(r["Review status"] in gtc.HUMAN_DECISIONS for r in log),
    }
    if freeze_manifest and freeze_manifest.exists():
        manifest = json.loads(freeze_manifest.read_text())
        checks["freeze_status"], checks["freeze_blocking"] = manifest["status"], manifest["blocking"]
    return checks


def write_workbook(path: Path, rows: list[dict], auto: list[dict], checks: dict, meta: dict) -> None:
    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "Read me"
    log_ws = wb.create_sheet("Change Log")
    log_ws.append(COLUMNS)
    for cell in log_ws[1]:
        cell.fill, cell.font = PatternFill("solid", fgColor="1F4E78"), Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    colours = (("UNCERTAIN", "F8CBAD"), ("P1 violation", "FFE699"), ("HIGH IMPACT", "DDEBF7"))
    for row in rows:
        log_ws.append([row[c] for c in COLUMNS])
        for key, colour in colours:
            if key in row["Flags"]:
                log_ws.cell(log_ws.max_row, COLUMNS.index("Flags") + 1).fill = PatternFill("solid", fgColor=colour)
                break
    grey = PatternFill("solid", fgColor="EDEDED")
    for r in auto:
        informational = r["Review status"] == "Informational"
        log_ws.append([r["Change ID"], r["Incident ID"], "Informational (no change)" if informational
                       else "Mechanical (auto-verified)", r["Sheet"], r["Record ID"], r["Row"], r["Field"], r["Rule"],
                       r["Original value"], r["Proposed value"], None, None, r["Independent verification"],
                       "Recorded only; nothing to decide" if informational else "Already applied; reject only if wrong",
                       "INFORMATIONAL" if informational else "AUTO-VERIFIED",
                       r["Review status"], r.get("Reviewer"), r.get("Review notes"), r.get("Accepted value"), None,
                       None, None])
        for cell in log_ws[log_ws.max_row]:
            cell.fill = grey
    status = DataValidation(type="list", formula1='"Pending,Approved,Rejected,Auto-verified,Informational"',
                            allow_blank=False)
    log_ws.add_data_validation(status)
    status.add(f"P2:P{max(log_ws.max_row, 2)}")
    for i, width in enumerate([10, 14, 20, 13, 9, 6, 16, 6, 50, 50, 13, 9, 50, 26, 30, 12, 12, 26, 40, 40, 60, 24],
                              start=1):
        log_ws.column_dimensions[get_column_letter(i)].width = width
    for row in log_ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    log_ws.freeze_panes = "C2"
    log_ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{log_ws.max_row}"

    by_incident = wb.create_sheet("By incident")
    by_incident.append(["Incident ID", "Pending", "Types", "Uncertain", "P1 violations", "Pronoun checks",
                        "High impact", "Frames"])
    groups: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        groups[row["Incident ID"]].append(row)
    for incident, items in sorted(groups.items()):
        by_incident.append([incident, sum(i["Review status"] == "Pending" for i in items),
                            ", ".join(sorted({i["Correction type"] for i in items})),
                            sum("UNCERTAIN" in i["Flags"] for i in items),
                            sum("P1 violation" in i["Flags"] for i in items),
                            sum("permitted under P1" in i["Flags"] for i in items),
                            sum("HIGH IMPACT" in i["Flags"] for i in items), f"eval_data/gt_frames/{incident}/"])
    for cell in by_incident[1]:
        cell.font = Font(bold=True)

    version = meta["version"]
    lines = [
        (f"Ground Truth Review Package {version}", True),
        (f"Human decisions on the Pending entries of {meta['candidate_name']} ({meta['candidate_sha256'][:16]}…). "
         "A new file: the candidate, the earlier review workbooks and the source are never edited. Generated by "
         f"{GENERATOR}; nothing here is approved until you set it.", False),
        ("", False), ("How to review", True),
        ("1. Before and after every session (Excel closed), from eval/: uv run python scripts/gt_review_backup.py "
         f"snapshot --kind Review --version 'Package {version}' --reason 'before session 1'.", False),
        ("2. Work in 'Change Log', one incident at a time. For UNCERTAIN and HIGH IMPACT rows watch the local clip "
         "at the cited seconds (frames: column Frames, or the HTML companion).", False),
        ("3. Review status: Approved = apply the AI proposal, or your 'Accepted value' if filled in; on a 'no "
         "change' row it confirms the original. Rejected = keep the original. Leave Pending to come back later.",
         False),
        ("4. To use a suggestion, copy 'Suggested accepted value' into 'Accepted value'. Never edit Before or "
         "After. Out-of-taxonomy values are exact JSON, e.g. {\"Out_Of_Taxonomy\": true, \"Nearest_Type\": "
         "\"explosion\"}.", False),
        ("5. Fill Reviewer on every row you decide; add Review notes for any edit or rejection.", False),
        ("6. Leave the grey AUTO-VERIFIED rows unless a mechanical fix is wrong (then Rejected, with the reason).",
         False),
        ("7. Decide the two halves of a reclassification (R10) the same way.", False),
        ("8. Then: uv run python scripts/gt_candidate_build.py --prior \"ground_truth/review/Ground Truth Review "
         f"Package {version}.xlsx\" --version v0.3   (a new candidate; refuses if this package does not match the "
         "audit).", False),
        ("9. Then: uv run python scripts/gt_freeze.py --candidate \"ground_truth/review/Ground Truth Candidate "
         "v0.3.xlsx\" shows what still blocks the freeze; freezing needs your approval record.", False),
        ("", False), ("Pronoun policy P1 (adopted 2026-10-27)", True),
        ("Gendered pronouns are permitted after a visually supported, qualified gender identification (e.g. "
         "'appears to be a man'). Where gender cannot be determined and the description is gender-neutral, "
         "later pronouns must be neutral too. 'P1 violation' rows break this and carry a suggested fix; "
         "'permitted under P1' rows are fine if you agree the qualified identification is visually supported.",
         False),
        ("", False), ("Traceability checks (generated)", True),
    ] + [(f"{k}: {v}", False) for k, v in checks.items()]
    for text, bold in lines:
        readme.append([text])
        readme.cell(readme.max_row, 1).font = Font(bold=bold, size=13 if bold else 11)
        readme.cell(readme.max_row, 1).alignment = Alignment(wrap_text=True)
    readme.column_dimensions["A"].width = 140
    wb.save(path)


def _tile(frames_dir: Path, incident: str, sheet: str, t: float) -> str | None:
    from PIL import Image

    frames_file = frames_dir / incident / "frames.json"
    if not frames_file.exists():
        return None
    entry = next((s for s in json.loads(frames_file.read_text())["sheets"] if s["file"] == sheet), None)
    if entry is None or not entry["times"] or not (frames_dir / incident / sheet).exists():
        return None
    index = min(range(len(entry["times"])), key=lambda i: abs(entry["times"][i] - t))
    image = Image.open(frames_dir / incident / sheet)
    columns, width = 6, 320
    tile_height = image.height // ((len(entry["times"]) + columns - 1) // columns)
    left, top = (index % columns) * width, (index // columns) * tile_height
    buffer = io.BytesIO()
    image.crop((left, top, left + width, top + tile_height)).convert("RGB").save(buffer, format="JPEG", quality=70)
    return base64.b64encode(buffer.getvalue()).decode()


def write_html(path: Path, rows: list[dict], version: str, frames_dir: Path) -> None:
    parts = [f"<!doctype html><meta charset=utf-8><title>GT Review Package {html.escape(version)}</title><style>"
             "body{font:14px/1.45 system-ui,sans-serif;margin:16px;max-width:1200px;background:#fff;color:#111}"
             "h2{border-top:2px solid #1f4e78;padding-top:8px}.item{border:1px solid #ccc;border-radius:6px;"
             "padding:8px 12px;margin:10px 0}.flag{display:inline-block;background:#eee;border-radius:4px;"
             "padding:1px 6px;margin-right:4px;font-size:12px}.grid{display:grid;grid-template-columns:"
             "minmax(0,1fr) minmax(0,1fr);gap:8px}.box{background:#f7f7f7;padding:6px;border-radius:4px;"
             "white-space:pre-wrap;word-break:break-word}img{width:240px;margin:4px 4px 0 0;border:1px solid #999}"
             ".rec{font-weight:600;color:#1f4e78}@media(max-width:700px){.grid{grid-template-columns:minmax(0,1fr)}}"
             "</style>", f"<h1>Ground Truth Review Package {html.escape(version)}</h1><p>{len(rows)} entries, grouped "
             "by incident and type. Decisions go in the workbook, not here. Local and private: annotation text.</p>"]
    groups: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        groups[row["Incident ID"]].append(row)
    for incident, items in sorted(groups.items()):
        parts.append(f"<h2>{html.escape(incident)} ({len(items)})</h2>")
        for row in items:
            flags = "".join(f"<span class=flag>{html.escape(f)}</span>" for f in row["Flags"].split("; ") if f)
            images = []
            for evidence in row["_evidence"][:4]:
                data = _tile(frames_dir, incident, evidence["sheet"], float(evidence["t"]))
                if data:
                    images.append(f"<figure style='display:inline-block;margin:0'><img alt='' src='data:image/jpeg;"
                                  f"base64,{data}'><figcaption>{evidence['t']} s: "
                                  f"{html.escape(evidence.get('observation', ''))}</figcaption></figure>")
            suggestion = (f"<p><b>Suggested accepted value:</b></p><div class=box>"
                          f"{html.escape(str(row['Suggested accepted value']))}</div>"
                          if row["Suggested accepted value"] else "")
            parts.append(
                f"<div class=item><p><b>{row['Change ID']}</b> · {html.escape(row['Correction type'])} · "
                f"{html.escape(str(row['Record ID']))} · {html.escape(str(row['Field']))} · AI: {row['AI decision']} "
                f"({row['AI confidence']}) · status: {html.escape(str(row['Review status']))} {flags}</p><div "
                f"class=grid><div><b>Before</b><div class=box>{html.escape(str(row['Before']))}</div></div><div>"
                f"<b>After</b><div class=box>{html.escape(str(row['After (AI proposal)']))}</div></div></div>"
                f"{suggestion}<p class=rec>Recommendation: {html.escape(row['Recommendation'])}</p>"
                f"<p>{html.escape(row['Why'])}</p><div>{''.join(images) or '<i>no evidence tiles</i>'}</div></div>")
    path.write_text("".join(parts))


def build(candidate_path: Path, version: str, *, source: Path = DEFAULT_SOURCE, proposals_path: Path | None = None,
          out_dir: Path = DEFAULT_OUT_DIR, html_dir: Path = HTML_DIR, frames_dir: Path = FRAMES_DIR,
          prior_package: Path | None = None, overrides_path: Path | None = None,
          freeze_manifest: Path | None = None) -> dict:
    candidate_version = candidate_path.stem.rsplit(" ", 1)[-1]
    proposals_path = proposals_path or candidate_path.parent / f"ai-proposals-{candidate_version}.jsonl"
    overrides_path = overrides_path or candidate_path.parent / f"review-overrides-{candidate_version}.json"
    output = out_dir / f"Ground Truth Review Package {version}.xlsx"
    inputs = [p for p in (source, candidate_path, proposals_path, prior_package) if p]
    gtc.guard_inputs(inputs)
    if output.exists():
        raise review_backup.OverwriteRefusedError(f"{output.name} exists; build a new --version (never overwritten)")
    if output.resolve() in {p.resolve() for p in inputs}:
        raise review_backup.OverwriteRefusedError("the package must be a new file, not an input")
    before = {p: sha256_file(p) for p in inputs}
    candidate = openpyxl.load_workbook(candidate_path, data_only=True)
    log = gtc.read_change_log(candidate)
    proposals = gtc.load_proposals(proposals_path)
    overrides = json.loads(overrides_path.read_text()) if overrides_path.exists() else {}
    prior: dict[str, dict] = {}
    if prior_package is not None:
        prior_log = gtc.read_change_log(openpyxl.load_workbook(prior_package, data_only=True))
        if {str(r["Change ID"]) for r in prior_log} != {str(r["Change ID"]) for r in log}:
            raise ValueError(f"{prior_package.name} covers different change ids than {candidate_path.name}")
        prior = {str(r["Change ID"]): {k: r.get(k) for k in HUMAN_COLUMNS} for r in prior_log
                 if r.get("Review status") in gtc.HUMAN_DECISIONS or r.get("Reviewer") or r.get("Review notes")}
    rows = build_rows(log, proposals, overrides, prior)
    auto = [{**r, **prior.get(str(r["Change ID"]), {})} for r in log
            if r["Review status"] in ("Auto-verified", "Informational")]
    checks = verification(source, candidate, log, proposals, freeze_manifest)
    out_dir.mkdir(parents=True, exist_ok=True)
    write_workbook(output, rows, auto, checks, {"version": version, "candidate_name": candidate_path.name,
                                                "candidate_sha256": before[candidate_path]})
    html_dir.mkdir(parents=True, exist_ok=True)
    write_html(html_dir / f"review-package-{version}.html", rows, version, frames_dir)
    checks["inputs_unchanged"] = before == {p: sha256_file(p) for p in inputs}
    (html_dir / f"verification-{version}.json").write_text(json.dumps(checks, indent=2) + "\n")
    if not checks["inputs_unchanged"]:
        raise RuntimeError("an input changed while the package was built")
    return {"output": str(output), "output_sha256": sha256_file(output), "rows": len(rows),
            "carried_decisions": len(prior), "by_type": dict(Counter(r["Correction type"] for r in rows)),
            "flags": dict(Counter(f.split(" (")[0] for r in rows for f in r["Flags"].split("; ") if f)),
            "statuses": dict(Counter(r["Review status"] for r in rows)), "checks": checks}


def main() -> int:
    parser = argparse.ArgumentParser(description="Build the ground-truth human review package.")
    parser.add_argument("--candidate", type=Path, default=DEFAULT_OUT_DIR / "Ground Truth Candidate v0.2.xlsx")
    parser.add_argument("--version", required=True, help="Package version, e.g. v0.2.1 (a new file every time)")
    parser.add_argument("--prior-package", type=Path, help="An earlier package whose decisions carry forward")
    parser.add_argument("--overrides", type=Path, help="Local overrides JSON (default review-overrides-<v>.json)")
    parser.add_argument("--proposals", type=Path, help="AI proposals JSONL (default ai-proposals-<candidate v>.jsonl)")
    args = parser.parse_args()
    try:
        result = build(args.candidate, args.version, prior_package=args.prior_package, overrides_path=args.overrides,
                       proposals_path=args.proposals,
                       freeze_manifest=EVAL_DIR / "eval_data" / "ground_truth" / "ground-truth-v1.0" / "MANIFEST.json")
    except review_backup.OverwriteRefusedError as exc:
        print(f"REFUSED: {exc}", file=sys.stderr)
        return 2
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
